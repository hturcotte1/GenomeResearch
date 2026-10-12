"""Readers for the raw files that feed the namespace and the merged table.

Each reader returns names exactly as written in the raw file; no
normalisation happens here (that is ``src/namespace/mapping.py``). Every
reader reads only from ``data/raw/<source>/`` and is deterministic.

Orientation conventions, as established and verified in Phase 1:

- Randi atlas matrices are [responding, stimulated].
- NemaNode records, Fenyves sheet 5, Ripoll-Sánchez matrices and Bentley edge
  lists are (pre/sender, post/receiver).
"""

from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass
from pathlib import Path

from src.acquire.common import RAW_ROOT


def _raw(source: str, name: str) -> Path:
    return RAW_ROOT / source / name


@dataclass(frozen=True)
class Edge:
    pre: str
    post: str
    kind: str  # "chemical" | "electrical" | free text for extrasynaptic layers
    weight: float
    detail: str = ""


# ---------------------------------------------------------------- Randi 2023

def randi_atlas(filename: str) -> dict:
    """Load one atlas release. Matrices are [responding, stimulated].

    Returns {"ids": [...], "time_compiled": str, "<strain>": {"q", "q_eq",
    "dFF", "occ1", "n_trials", "kernels"}}. ``n_trials`` is len(dFF_all), the
    number of retained trials. ``kernels`` stays an object array of flat
    float arrays (g, factor, power_t, branch quadruples).
    """
    import h5py
    import numpy as np

    out: dict = {}
    with h5py.File(_raw("randi2023", filename), "r") as f:
        out["ids"] = [x.decode("utf-8") for x in f["neuron_ids"][:]]
        out["time_compiled"] = f.attrs["time_compiled"].decode("utf-8")
        for s in ("wt", "unc31"):
            g = f[s]
            out[s] = {
                "q": g["q"][:],
                "q_eq": g["q_eq"][:],
                "dFF": g["dFF"][:],
                "occ1": g["occ1"][:],
                "n_trials": np.vectorize(len, otypes=[int])(g["dFF_all"][:]),
                "kernels": g["kernels"][:],
            }
    return out


def randi_ganglia() -> dict:
    return json.loads(_raw("randi2023", "aconnectome_ids_ganglia.json").read_text(encoding="utf-8"))


# ---------------------------------------------------------------- NemaNode-format anatomy

_NEMANODE_TYP = {0: "chemical", 2: "electrical"}


def nemanode_tsv(source: str, filename: str) -> list[Edge]:
    """NemaNode CSV download (tab-separated pre/post/type/synapses)."""
    text = _raw(source, filename).read_bytes().decode("ascii").replace("\r\n", "\n")
    lines = [ln for ln in text.split("\n") if ln != ""]
    header = lines[0].split()
    if header != ["pre", "post", "type", "synapses"]:
        raise ValueError(f"{filename}: unexpected header {header}")
    edges = []
    for ln in lines[1:]:
        pre, post, kind, w = ln.split("\t")
        edges.append(Edge(pre, post, kind, float(int(w))))
    return edges


def nemanode_json(source: str, filename: str) -> list[Edge]:
    """NemaNode source records. Weight = number of synapse records (len(syn)).

    This is the unit of the Witvliet et al. paper's synapse-count tables
    (verified in Phase 1). Electrical records appear in both directions.
    """
    entries = json.loads(_raw(source, filename).read_text(encoding="utf-8"))
    out = []
    for e in entries:
        kind = _NEMANODE_TYP.get(e["typ"])
        if kind is None:
            raise ValueError(f"{filename}: unknown typ {e['typ']}")
        out.append(Edge(e["pre"], e["post"], kind, float(len(e["syn"]))))
    return out


# ---------------------------------------------------------------- Fenyves 2020

FENYVES_XLSX = "journal.pcbi.1007974.s003.xlsx"
_FEN_CLASSES = ("Glu+", "Glu-", "ACh+", "ACh-", "GABA+", "GABA-")


def fenyves_sign_rows() -> list[dict]:
    """Rows of sheet '5. Sign prediction' (journal version), in file order."""
    import openpyxl

    wb = openpyxl.load_workbook(_raw("fenyves2020", FENYVES_XLSX), read_only=True, data_only=True)
    ws = wb["5. Sign prediction"]
    rows = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r[0] is None:
            continue
        rows.append(
            {
                "pre": str(r[0]),
                "pre_nt_primary": r[1] if r[1] not in (None, 0) else "",
                "pre_nt_secondary": r[2] if r[2] not in (None, 0) else "",
                "post": str(r[3]),
                "synapses": int(r[4]),
                "edge_type": str(r[5]),
                "matched_classes": [c for c, flag in zip(_FEN_CLASSES, r[6:12]) if flag is True],
                "prediction": str(r[16]),
            }
        )
    wb.close()
    return rows


def fenyves_receptor_genes() -> tuple[dict[str, list[str]], dict[str, list[str]]]:
    """(class -> receptor genes) from sheet 2 and (neuron -> expressed receptor genes) from sheet 3."""
    import openpyxl

    wb = openpyxl.load_workbook(_raw("fenyves2020", FENYVES_XLSX), read_only=True, data_only=True)
    classes: dict[str, list[str]] = {c: [] for c in _FEN_CLASSES}
    for r in wb["2. Receptor gene table"].iter_rows(min_row=2, values_only=True):
        for c, g in zip(_FEN_CLASSES, r[:6]):
            if g:
                classes[c].append(str(g))
    expressed: dict[str, list[str]] = {}
    for r in wb["3. Receptor Gene expression"].iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        expressed[str(r[0])] = sorted(str(g) for g in r[1:] if g)
    wb.close()
    return {c: sorted(v) for c, v in classes.items()}, expressed


# ---------------------------------------------------------------- Ripoll-Sánchez 2023

def ripoll_matrix(filename: str) -> list[Edge]:
    """Header-row matrix CSV (row = sender, column = receiver). Nonzero cells only."""
    rows = list(csv.reader(io.StringIO(_raw("ripollsanchez2023", filename).read_text(encoding="utf-8"))))
    cols = rows[0][1:]
    out = []
    for r in rows[1:]:
        pre = r[0]
        for post, v in zip(cols, r[1:]):
            w = float(v)
            if w != int(w):
                raise ValueError(f"{filename}: non-integer weight {v}")
            if w:
                out.append(Edge(pre, post, "neuropeptide", w))
    return out


def ripoll_labels(filename: str) -> tuple[list[str], list[str]]:
    rows = list(csv.reader(io.StringIO(_raw("ripollsanchez2023", filename).read_text(encoding="utf-8"))))
    return rows[0][1:], [r[0] for r in rows[1:]]


# ---------------------------------------------------------------- Bentley 2016

def bentley_edgelist(filename: str) -> list[Edge]:
    """source,target,ligand,receptor rows (no header); kind = ligand, detail = receptor."""
    text = _raw("bentley2016", filename).read_bytes().decode("ascii")
    out = []
    for ln in text.split("\n"):
        if not ln:
            continue
        f = ln.split(",")
        out.append(Edge(f[0], f[1], f[2], 1.0, f[3]))
    return out


# ---------------------------------------------------------------- CeNGEN (Taylor 2021)

def cengen_types(filename: str = "cengen_021821_medium_threshold2.csv") -> list[str]:
    with open(_raw("taylor2021", filename), encoding="utf-8") as fh:
        header = next(csv.reader(fh))
    return header[3:]
