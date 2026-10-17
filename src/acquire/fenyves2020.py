"""Fenyves et al. 2020, PLoS Comput Biol 16(12):e1007974: predicted synaptic polarity.

Fenyves et al. predicted the sign of each ionotropic chemical connection from
presynaptic neurotransmitter (NT) and postsynaptic ionotropic receptor (R)
gene expression ("NT+R" method). Every machine-readable copy that is
reachable from this environment is acquired here; all are SECONDARY copies.
See docs/provenance/fenyves2020.md.

Files and what they are:

1. ``journal.pcbi.1007974.s003.xlsx``: the PLOS supporting-information item
   ``.s003``, as bundled in the ``wormneuroatlas`` 0.0.7.3 wheel (Randi et al.).
   Sheets: USER GUIDE, 1. NT expr, 2. Receptor gene table, 3. Receptor Gene
   expression, 4. Receptor type expression (trailing space in the sheet name),
   5. Sign prediction, Summary. Byte-identical to the copy in GitHub
   francescorandi/wormneuroatlas @ b2e13d88. Randi et al.'s analysis code
   (leiferlab/pumpprobe) downloads the same DOI (``.s003``) and reads sheet 5.
2. Authors' GitHub repository bank-fenyves/CeConn-SignPrediction @ e657d5f:
   - ``Data (rev 12-09-2020).csv`` (uploaded 2020-09-14): sheet 5 columns A..Q
     as ';'-separated text. Cell-for-cell equal to the xlsx sheet 5 values.
   - ``data.csv`` (uploaded 2020-01-23): an EARLIER version. Same 3,638
     chemical rows plus 2,167 electrical rows and one legend row; booleans in
     Hungarian (IGAZ/HAMIS). 13 chemical rows carry a different prediction
     (ACh inputs to ADEL/ADER/AIYL/AIYR). Its totals equal the bioRxiv
     preprint's (9,034 / 2,580 / 3,431), not the journal's.
   - ``SignPred Calc``, ``Search Tool`` (R scripts) and ``README.md``.
3. Hermaphrodite edge lists (Source, Target, Weight, Type), all OpenWorm
   copies of a list whose WormWiring origin is presumed but UNVERIFIED:
   - ``herm_full_edgelist_CElegansNeuroML-3808b6b.csv``: the earliest copy
     found, openworm/CElegansNeuroML @ 3808b6b (2016-12-06, blob a1313b0e).
     Its neuron-to-neuron rows equal ``data.csv``'s chemical AND electrical
     rows (pairs and weights).
   - ``herm_full_edgelist.csv`` (pinned ConnectomeToolbox commit, blob
     89ffc2a2): the same list plus one row OpenWorm added the same day
     (CElegansNeuroML 1ca08b1, "Add missing GJ PVCR-PVDR"): PVCR -> PVDR,
     2, electrical. Its chemical rows whose two endpoints are both in the
     xlsx's own 302-name neuron list are exactly sheet 5's 3,638 (source,
     target, weight) rows.
   - ``herm_full_edgelist_MODIFIED.csv`` (pinned ConnectomeToolbox commit,
     blob 5065baf7): a later OpenWorm manual edit. Kept as evidence that it
     does NOT match sheet 5.
   None of them is Cook et al. 2019 SI 5 (see
   scripts/provenance_checks/fenyves2020_vs_cook2019_si5.py).

Facts later phases must respect (checked by ``record_counts``):

- Sheet 5 rows are directed chemical connections, Source (col A) ->
  Target (col D), weight = synapse count (col E). Data rows are 3..3640;
  rows 1-2 are a two-line header. Self-pairs occur (34 rows).
- Names: single-digit ventral-cord indices are zero-padded (VA01, AS01);
  two-digit ones are not (AS10). AWC as AWCL/AWCR; CANL/CANR are in sheet 1
  but in no connection. In the edge lists every name is padded with spaces.
- Workbook values are Excel's cached formula results. Excel's own named
  ranges ``herm_full_edgelist``/``pred_polarity``/``synapses``/... stop at
  row 3477 (163 rows short) and are stale; the Summary sheet does not use
  them (it sums whole columns).
"""

from __future__ import annotations

import csv
import io
import re
import warnings
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from urllib.parse import quote

from .common import RAW_ROOT, ArchiveMember, RemoteFile, acquire_source, verify_pypi_digest
from .distributions import (
    CONNECTOME_TOOLBOX_COMMIT,
    WORMNEUROATLAS_PACKAGE,
    WORMNEUROATLAS_VERSION,
    WORMNEUROATLAS_WHEEL,
    connectome_toolbox_file,
)

SOURCE_ID = "fenyves2020"
CITATION = (
    "Fenyves BG, Szilágyi GS, Vassy Z, Sőti C, Csermely P. Synaptic polarity and sign-balance "
    "prediction using gene expression data in the Caenorhabditis elegans chemical synapse neuronal "
    "connectome network. PLoS Comput Biol 16(12): e1007974 (2020). doi:10.1371/journal.pcbi.1007974"
)

PRIMARY_S003 = (
    "https://doi.org/10.1371/journal.pcbi.1007974.s003 (PLOS Comput Biol supporting-information item "
    "s003 of doi:10.1371/journal.pcbi.1007974; its caption, inferred to be 'S1 Data' (NT+R prediction), "
    "is UNVERIFIED; see docs/provenance/fenyves2020.md)"
)

# The first author's own repository, cited in its README as materials of the paper.
SIGNPRED_REPO = "bank-fenyves/CeConn-SignPrediction"
SIGNPRED_COMMIT = "e657d5f3c3986e99ccd39601b245b9b32c22d949"  # 2021-01-07, repository HEAD
_SIGNPRED_PRIMARY = (
    f"https://github.com/{SIGNPRED_REPO} (the first author's own repository; no other copy known). "
    "Journal counterpart for the prediction table: " + PRIMARY_S003
)

_WORMWIRING_PRIMARY = (
    "Origin UNVERIFIED. The earliest copy found is openworm/CElegansNeuroML commit "
    "3808b6b235aa39ffac48bee319b06598bd23ff4d (2016-12-06, 'Add herm_full_edgelist.csv', blob a1313b0e), "
    "whose commit names no source. Presumed origin: WormWiring (Emmons lab) data downloads, "
    "https://wormwiring.org/series/ ('Edge list (.csv)' for the adult hermaphrodite), as the xlsx USER GUIDE "
    "says 'Source data: Wormwiring connectome reconstruction'; no WormWiring-hosted copy was compared, and "
    "its exact file URL and release date are UNVERIFIED."
)

# OpenWorm's CElegansNeuroML repository: the earliest copy of the edge list found.
CENML_REPO = "openworm/CElegansNeuroML"
CENML_PRE_EDIT_COMMIT = "3808b6b235aa39ffac48bee319b06598bd23ff4d"  # 2016-12-06 10:51:32 UTC, lungd


def _signpred(path: str, sha256: str, size: int, blob: str, note: str) -> RemoteFile:
    return RemoteFile(
        url=f"https://raw.githubusercontent.com/{SIGNPRED_REPO}/{SIGNPRED_COMMIT}/{quote(path)}",
        filename=path,
        sha256=sha256,
        size=size,
        distributor=f"GitHub {SIGNPRED_REPO} @ {SIGNPRED_COMMIT} (B. G. Fenyves, first author)",
        origin="secondary",
        primary_reference=_SIGNPRED_PRIMARY,
        git_blob_sha1=blob,
        note=note,
    )


XLSX = "journal.pcbi.1007974.s003.xlsx"
REV_CSV = "Data (rev 12-09-2020).csv"
JAN_CSV = "data.csv"
EDGELIST = "herm_full_edgelist.csv"
EDGELIST_PRE_EDIT = "herm_full_edgelist_CElegansNeuroML-3808b6b.csv"
EDGELIST_MODIFIED = "herm_full_edgelist_MODIFIED.csv"

FILES = (
    ArchiveMember(
        archive=WORMNEUROATLAS_WHEEL,
        member=f"wormneuroatlas/data/{XLSX}",
        filename=XLSX,
        sha256="85959066fd7cbdbc2024d0ebb323b71c4365f4083bc85e555ee973f470697c47",
        size=983199,
        primary_reference=PRIMARY_S003,
        note=(
            "NT+R sign-prediction workbook (USER GUIDE: 'Source data: Wormwiring connectome reconstruction'; "
            "'Method of prediction: NT+R'). docProps: created 2019-07-01T09:46:44Z, modified "
            "2020-09-10T23:15:24Z, lastModifiedBy 'Fenyves Bánk'. wormneuroatlas SynapseSign reads sheets 1-2 "
            "only; NeuroAtlas.get_chemical_synapse_sign recomputes signs from CeNGEN. pumpprobe Funatlas reads "
            "sheet 5 from doi .s003. Byte-identical to GitHub francescorandi/wormneuroatlas@b2e13d88 blob "
            "0457b2faf3d051c1826b9865421c258f2dc37509. The wormneuroatlas wheels 0.0.6, 0.0.7.1 and 0.0.7.2 carry a "
            "byte-identical member (same SHA-256 and size); 0.0.2 to 0.0.5 do not contain the file."
        ),
    ),
    _signpred(
        REV_CSV,
        "8748f968d759157e9e955b325d59314cc0ad8f4bf4501521f380a73a674c66ff",
        321820,
        "660e2100158cc40bf37262075dba59c82ed4d1e1",
        "Uploaded 2020-09-14 (commit b2ae4a9). 3,638 chemical rows, 17 columns = xlsx sheet 5 A..Q, "
        "cell-for-cell equal to the xlsx values. ';'-separated, UTF-8 with BOM, CRLF, no final newline.",
    ),
    _signpred(
        JAN_CSV,
        "0d45dcf0970b94a90f23eef51695f9a47a69bbef9615920915c824877959c9f7",
        531133,
        "2442ce8f225cdc82c8d8f3363f38f83ea9c2f1fb",
        "Uploaded 2020-01-23 (commit 623f4fb), never modified. EARLIER prediction version: 3,638 chemical + "
        "2,167 electrical rows + 1 legend row; booleans IGAZ/HAMIS. 13 chemical predictions differ from the "
        "xlsx. ';'-separated, no BOM, CRLF, no final newline. The repository README still describes this file.",
    ),
    _signpred(
        "SignPred Calc",
        "0966414de2c990ecd70747127797d00e7f54e8eff808ad345a6a4edb895adb8b",
        2509,
        "98f566a3632b81f82484cb945eec1bc57462c99a",
        "R script (authors' code): sums Synapses by Pred_pol over Type == 'chemical' rows of data.csv. "
        "Kept as documentation of the authors' count definitions; never executed.",
    ),
    _signpred(
        "Search Tool",
        "94641ae0887ea955477ae348f40ccf3e0fa759532d3e66aba4c8db6449b49d15",
        465,
        "ebd78dfe12369207cabbad276b09720a39ad3595",
        "R script (authors' code): interactive lookup of one connection. Documentation only; never executed.",
    ),
    _signpred(
        "README.md",
        "beff99dfcb1b79f34b56b11c6e83c19291ccff20a61c30be3ebab693c52c0142",
        968,
        "ce6c7c596f601f2bdb937dd746b3bc0bf479d849",
        "The authors' repository README: links the repository to the paper; describes data.csv only.",
    ),
    connectome_toolbox_file(
        f"cect/data/{EDGELIST}",
        sha256="142693f17556148d7f962835b18ac6dd5af18b7467eef61815ebc1dd5474c0ca",
        size=245463,
        git_blob_sha1="89ffc2a2d9e1934ddc462d51032acead5e4cfc91",
        primary_reference=(
            "OpenWorm-amended copy: openworm/CElegansNeuroML 1ca08b14f4e795a973ce447565e0c97b897ee4dc "
            "(2016-12-06, 'Add missing GJ PVCR-PVDR to herm_full_edgelist.csv') added the one row PVCR -> PVDR, "
            "weight 2, electrical (written space-padded like every row: 'PVCR , PVDR        ,2,electrical') to "
            "blob a1313b0e (acquired as " + EDGELIST_PRE_EDIT + "), giving this "
            "blob 89ffc2a2; the same blob is in CElegansNeuroML master and in openworm/c302 (49acae1, 2018-07-06). "
            + _WORMWIRING_PRIMARY
        ),
        note=(
            "OpenWorm copy of a hermaphrodite edge list (chemical + electrical, neurons + muscles/other) with one "
            "OpenWorm-added gap junction (PVCR->PVDR, 2); WormWiring origin UNVERIFIED. Its neuron-to-neuron "
            "chemical rows equal xlsx sheet 5 exactly. The toolbox labels it 'legacy dataset (based on Cook et al. "
            "2019) ... Do not use!' (cect/readers/UpdatedSpreadsheetDataReader.py @ "
            f"{CONNECTOME_TOOLBOX_COMMIT[:7]}). Not the Cook 2019 SI 5 matrix. Every name is space-padded."
        ),
    ),
    RemoteFile(
        url=f"https://raw.githubusercontent.com/{CENML_REPO}/{CENML_PRE_EDIT_COMMIT}/herm_full_edgelist.csv",
        filename=EDGELIST_PRE_EDIT,
        sha256="7dfc27e4fa04844cd8b38554da5f4765c4a90ddce313d0f2612ee03e526829a6",
        size=245430,
        distributor=f"GitHub {CENML_REPO} @ {CENML_PRE_EDIT_COMMIT} (path herm_full_edgelist.csv)",
        origin="secondary",
        primary_reference=_WORMWIRING_PRIMARY,
        git_blob_sha1="a1313b0e53d258c361032d08657ce97566066ced",
        note=(
            "Pre-edit OpenWorm edge list, as first committed (CElegansNeuroML 3808b6b, 2016-12-06, 'Add "
            "herm_full_edgelist.csv'). Differs from " + EDGELIST + " only by lacking the one row PVCR -> PVDR, "
            "weight 2, electrical (written space-padded like every row: 'PVCR , PVDR        ,2,electrical'). "
            "Its neuron-to-neuron rows equal data.csv's chemical and electrical rows (pairs and "
            "weights). Stored under a distinct local name; the repository path is herm_full_edgelist.csv. "
            "LF line ends, no final newline, every name space-padded."
        ),
    ),
    connectome_toolbox_file(
        f"cect/data/{EDGELIST_MODIFIED}",
        sha256="8ba817eaed196e4aff3c6646af3d9493f5929ca6cea134439ee1e181087ac205",
        size=245926,
        git_blob_sha1="5065baf77a58e990a28f460a1e4c1d6c5d2eecaf",
        primary_reference=(
            "No external primary: an OpenWorm manual edit of blob 89ffc2a2 (" + EDGELIST + "). Blob 5065baf7 "
            "first appears at path herm_full_edgelist.csv in openworm/CElegansNeuroML "
            "1c24a6ce1cc27051f6e9765e1aff837a540ae260 (2017-10-12, 'intermediate commit') and again in "
            "a40fccb764975bfc9dd13c19ea7e0dd190cc5186 (2018-01-26, 'DO NOT MERGE!! Manual changes to connectome!'), "
            "neither on CElegansNeuroML master; both commits' parent holds blob 89ffc2a2 there. The same blob is "
            "openworm/c302 c302/data/herm_full_edgelist_MODIFIED.csv, added in c302 "
            "3e723fb840f220268e7b5e3181a6ba5f945b5b93 (2018-07-10)."
        ),
        note=(
            "Acquired as evidence only: NOT a Fenyves input. Equals " + EDGELIST + " minus the chemical row "
            "VB04->VB05 (1) plus 16 appended chemical rows of weight 5 (DB01->DB02 .. DB06->DB07, VB01->VB02 .. "
            "VB10->VB11), so on the xlsx's 302 neurons it has 15 chemical pairs sheet 5 lacks and VB04->VB05 = 5 "
            "where sheet 5 has 1."
        ),
    ),
)

_WHEEL_FILES = (XLSX,)


def acquire() -> list[dict]:
    pypi_check = verify_pypi_digest(
        WORMNEUROATLAS_PACKAGE,
        WORMNEUROATLAS_VERSION,
        WORMNEUROATLAS_WHEEL.filename,
        WORMNEUROATLAS_WHEEL.sha256,
    )
    extra = {name: [f"archive {WORMNEUROATLAS_WHEEL.filename}: {pypi_check}"] for name in _WHEEL_FILES}
    return acquire_source(SOURCE_ID, FILES, extra)


# ---------------------------------------------------------------------------
# Record counts
# ---------------------------------------------------------------------------

_S5 = "5. Sign prediction"
_S4 = "4. Receptor type expression "
_CATS = ("+", "-", "complex", "no pred")
_NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def _json_value(v):
    """Make a cell value JSON-safe without changing it (ints stay ints, bools stay bools)."""
    if v is None or isinstance(v, (bool, int, float, str)):
        return v
    return str(v)


def _load_workbooks(path):
    import openpyxl

    values = openpyxl.load_workbook(path, read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
    return values, formulas


def _rows(wb, sheet, min_row=1):
    return list(wb[sheet].iter_rows(min_row=min_row, values_only=True))


def _formula_text(x):
    """The formula a formula-workbook cell holds, or None for a typed value.

    openpyxl returns an ordinary formula as a str starting with '=', and an
    array (CSE) formula as an ``ArrayFormula`` object whose ``.text`` is the
    formula. Both are formulas.
    """
    from openpyxl.worksheet.formula import ArrayFormula

    if isinstance(x, ArrayFormula):
        return x.text
    if isinstance(x, str) and x.startswith("="):
        return x
    return None


def _sheet_cells(values_wb, formula_wb, sheet, min_col=1):
    """{coordinate: [cached value, formula or None]} for every non-empty cell at column >= min_col."""
    from openpyxl.utils import get_column_letter

    vrows = _rows(values_wb, sheet)
    frows = _rows(formula_wb, sheet)
    out = {}
    for i, (vr, fr) in enumerate(zip(vrows, frows), 1):
        for j, v in enumerate(vr, 1):
            if j < min_col or v is None:
                continue
            f = fr[j - 1] if j - 1 < len(fr) else None
            out[f"{get_column_letter(j)}{i}"] = [_json_value(v), _formula_text(f)]
    return out


def _defined_names(path) -> dict:
    """Excel defined names, as written in xl/workbook.xml. Sheet-local ones are keyed 'name@sheetIndex'."""
    with zipfile.ZipFile(path) as zf:
        root = ET.fromstring(zf.read("xl/workbook.xml"))
    out = {}
    for dn in root.iterfind("m:definedNames/m:definedName", _NS):
        key = dn.get("name")
        if dn.get("localSheetId") is not None:
            key = f"{key}@{dn.get('localSheetId')}"
        out[key] = dn.text
    return dict(sorted(out.items()))


def _sheet_dimensions(path) -> dict:
    """{sheet name: the <dimension ref> Excel stored for it (its used range)}, from the raw XML."""
    rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    with zipfile.ZipFile(path) as zf:
        wb = ET.fromstring(zf.read("xl/workbook.xml"))
        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        target = {r.get("Id"): r.get("Target") for r in rels}
        out = {}
        for sh in wb.iterfind("m:sheets/m:sheet", _NS):
            part = "xl/" + target[sh.get(rel_ns)].lstrip("/").removeprefix("xl/")
            dim = ET.fromstring(zf.read(part)).find("m:dimension", _NS)
            out[sh.get("name")] = dim.get("ref") if dim is not None else None
    return dict(sorted(out.items()))


def _by_category(rows, cat_idx, weight_idx):
    """{category: {"connections": rows, "synapses": sum of weight}} for the four NT+R categories."""
    out = {}
    for c in _CATS:
        sub = [r for r in rows if r[cat_idx] == c]
        out[c] = {"connections": len(sub), "synapses": sum(int(r[weight_idx]) for r in sub)}
    other = sorted({str(r[cat_idx]) for r in rows} - set(_CATS))
    if other:
        out["_other_values"] = other
    return out


def _ratio(a, b):
    return a / b if b else None


def _nt_r_category(g):
    """Sheet-5 column Q rule, transcribed from the workbook's own formula (cols O, P, Q)."""
    G, H, I, J, K, L = g
    o = (G and J) or (G and L) or (I and H) or (I and L) or (K and H) or (K and J)
    p = (G and H) or (I and J) or (K and L)
    if o or p:
        return "complex"
    if G or I or K:
        return "+"
    if H or J or L:
        return "-"
    return "no pred"


def _xlsx_counts(out: dict, raw) -> tuple[list, set]:
    path = raw / XLSX
    vwb, fwb = _load_workbooks(path)
    props = vwb.properties
    out["xlsx_sheet_names"] = list(vwb.sheetnames)
    out["xlsx_doc_created_utc"] = props.created.isoformat() if props.created else None
    out["xlsx_doc_modified_utc"] = props.modified.isoformat() if props.modified else None
    out["xlsx_doc_last_modified_by"] = props.lastModifiedBy
    out["xlsx_defined_names"] = _defined_names(path)
    out["xlsx_sheet_dimension_refs"] = _sheet_dimensions(path)

    out["xlsx_user_guide_cells"] = _sheet_cells(vwb, fwb, "USER GUIDE")
    out["xlsx_summary_sheet_cells"] = _sheet_cells(vwb, fwb, "Summary")
    out["xlsx_sheet4_results_block_cells"] = _sheet_cells(vwb, fwb, _S4, min_col=21)

    # Sheet 1: neuron list with NTs.
    s1 = [r for r in _rows(vwb, "1. NT expr", 2) if r[0] is not None]
    s1_names = [r[0] for r in s1]
    out["sheet1_rows"] = len(s1)
    out["sheet1_distinct_names"] = len(set(s1_names))
    out["sheet1_dominant_nt_counts"] = dict(sorted(Counter(str(r[1]) for r in s1).items()))
    out["sheet1_alternative_nt_counts"] = dict(sorted(Counter(str(r[2]) for r in s1).items()))
    out["sheet1_names_matching_zero_padded_index"] = sum(1 for n in s1_names if re.search(r"[A-Z]0\d$", n))
    out["sheet1_names_with_two_digit_index_not_padded"] = sorted(n for n in s1_names if re.search(r"[A-Z][1-9]\d$", n))
    out["sheet1_awc_names"] = sorted(n for n in s1_names if n.startswith("AWC"))
    out["sheet1_can_names"] = sorted(n for n in s1_names if n.startswith("CAN"))

    # Sheet 2: receptor genes per functional class (columns A..F).
    s2 = _rows(vwb, "2. Receptor gene table")
    classes = [str(h) for h in s2[0][:6]]
    out["sheet2_genes_per_class"] = {
        c: sum(1 for r in s2[1:] if r[j] is not None) for j, c in enumerate(classes)
    }
    genes = [r[j] for r in s2[1:] for j in range(6) if r[j] is not None]
    out["sheet2_gene_entries"] = len(genes)
    out["sheet2_distinct_genes"] = len(set(genes))

    # Sheet 3: per-neuron ragged list of expressed receptor genes (no header row).
    s3_all = _rows(vwb, "3. Receptor Gene expression")
    s3 = [r for r in s3_all if r[0] is not None]
    out["sheet3_neuron_rows"] = len(s3)
    out["sheet3_gene_tokens_not_in_sheet2"] = sorted({v for r in s3 for v in r[1:] if v is not None} - set(genes))
    out["sheet3_tokens_with_countif_wildcard_chars"] = sum(
        1 for r in s3 for v in r[1:] if isinstance(v, str) and any(c in v for c in "*?~")
    )

    # Sheet 4: per-neuron counts per receptor class (B..G), one formula per cell.
    s4_rownums = [i for i, r in enumerate(_rows(vwb, _S4), 1) if i >= 3 and r[0] is not None]
    s4_all, s4f_all = _rows(vwb, _S4), _rows(fwb, _S4)
    s4 = [s4_all[i - 1] for i in s4_rownums]
    s4f = [s4f_all[i - 1] for i in s4_rownums]
    out["sheet4_first_last_neuron_row"] = [s4_rownums[0], s4_rownums[-1]]
    out["sheet4_neuron_rows_contiguous"] = s4_rownums == list(range(s4_rownums[0], s4_rownums[-1] + 1))
    out["sheet4_neuron_rows"] = len(s4)
    out["sheet4_same_neuron_order_as_sheet1"] = [r[0] for r in s4] == s1_names
    out["sheet4_multi_receptor_complexity_true"] = sum(1 for r in s4 if r[8] is True)
    # Column I: the formula of the first neuron row as written, and the rows whose formula is not that
    # same formula with its row references moved to the row's own number.
    i_first, i_formula = s4_rownums[0], _formula_text(s4f[0][8])
    out["sheet4_colI_first_row_formula"] = i_formula
    out["sheet4_colI_rows_not_matching_first_row_formula_pattern"] = sum(
        1 for rownum, r in zip(s4_rownums, s4f)
        if i_formula is None
        or _formula_text(r[8]) != re.sub(rf"(?<=[A-Z]){i_first}(?!\d)", str(rownum), i_formula)
    )
    out["sheet4_cells_hardcoded_not_formula_by_column"] = {
        col: sum(1 for r in s4f if _formula_text(r[j]) is None) for col, j in zip("BCDEFG", range(1, 7))
    }
    from openpyxl.worksheet.formula import ArrayFormula

    out["sheet4_array_formula_cells_by_column"] = {
        col: sum(1 for r in s4f if isinstance(r[j], ArrayFormula)) for col, j in zip("BCDEFG", range(1, 7))
    }
    # Re-evaluate B..G: each cell is =SUMPRODUCT(COUNTIF(<class range>,'3. Receptor Gene expression'!$B<k>:$EL<k>))
    # with k = (sheet-4 row) - 1. COUNTIF is a case-insensitive equality test here (no token holds a
    # wildcard character), and an empty criteria cell counts as 0, which matches no gene name.
    from openpyxl.utils import column_index_from_string

    s4_names = ("glupos", "gluneg", "achpos", "achneg", "gabapos", "gabaneg")
    s2_rows_all = _rows(vwb, "2. Receptor gene table")
    class_genes = {}
    for nm in s4_names:
        m = re.fullmatch(r"'2\. Receptor gene table'!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)", out["xlsx_defined_names"][nm])
        c = column_index_from_string(m.group(1)) - 1
        class_genes[nm] = [
            str(s2_rows_all[i - 1][c]).lower()
            for i in range(int(m.group(2)), int(m.group(4)) + 1)
            if i - 1 < len(s2_rows_all) and c < len(s2_rows_all[i - 1]) and s2_rows_all[i - 1][c] is not None
        ]
    last_s3_col = column_index_from_string("EL")
    formula_off_pattern = Counter()
    mismatch = Counter()
    for rownum, vr, fr in zip(s4_rownums, s4, s4f):
        k = rownum - 1  # sheet-3 row the formula points to
        tokens = [str(t).lower() for t in s3_all[k - 1][1:last_s3_col] if t is not None]
        for col, j, nm in zip("BCDEFG", range(1, 7), s4_names):
            expect = f"=SUMPRODUCT(COUNTIF({nm},'3. Receptor Gene expression'!$B{k}:$EL{k}))"
            if _formula_text(fr[j]) != expect:
                formula_off_pattern[col] += 1
            if vr[j] != sum(class_genes[nm].count(t) for t in tokens):
                mismatch[col] += 1
    out["sheet4_cols_B_to_G_formulas_not_matching_expected_pattern"] = {c: formula_off_pattern[c] for c in "BCDEFG"}
    out["sheet4_cols_B_to_G_cached_values_differing_from_sheets2_3_reevaluation"] = {c: mismatch[c] for c in "BCDEFG"}
    out["sheet4_rows_whose_sheet3_row_names_a_different_neuron"] = sum(
        1 for rownum, r in zip(s4_rownums, s4) if s3_all[rownum - 2][0] != r[0]
    )

    # Sheet 5: the prediction table.
    s5_all = _rows(vwb, _S5)
    out["sheet5_header_row1"] = [_json_value(v) for v in s5_all[0]]
    out["sheet5_header_row2"] = [_json_value(v) for v in s5_all[1]]
    data_idx = [i for i, r in enumerate(s5_all, 1) if i >= 3 and r[0] is not None]
    nonempty_idx = [i for i, r in enumerate(s5_all, 1) if any(v is not None for v in r)]
    s5 = [s5_all[i - 1] for i in data_idx]
    # The formulas of the first data row (the anchor of each column's shared formula), as written.
    from openpyxl.utils import get_column_letter

    f_row3 = next(fwb[_S5].iter_rows(min_row=data_idx[0], max_row=data_idx[0], values_only=True))
    out["sheet5_first_data_row_formulas"] = {
        get_column_letter(j): _formula_text(v) for j, v in enumerate(f_row3, 1) if _formula_text(v) is not None
    }
    out["sheet5_data_rows"] = len(s5)
    out["sheet5_first_data_row"] = data_idx[0]
    out["sheet5_last_data_row"] = data_idx[-1]
    out["sheet5_data_rows_contiguous"] = data_idx == list(range(data_idx[0], data_idx[-1] + 1))
    out["sheet5_last_nonempty_row"] = nonempty_idx[-1]
    out["sheet5_rows_by_edge_type"] = dict(sorted(Counter(str(r[5]) for r in s5).items()))
    out["sheet5_edge_weight_sum_by_edge_type"] = {
        t: sum(int(r[4]) for r in s5 if str(r[5]) == t) for t in sorted({str(r[5]) for r in s5})
    }
    out["sheet5_edge_weight_sum"] = sum(int(r[4]) for r in s5)
    out["sheet5_edge_weight_min"] = min(int(r[4]) for r in s5)
    out["sheet5_edge_weight_max"] = max(int(r[4]) for r in s5)
    pairs = [(r[0], r[3]) for r in s5]
    out["sheet5_distinct_ordered_pairs"] = len(set(pairs))
    out["sheet5_sorted_by_source_then_target"] = pairs == sorted(pairs)
    selfp = [r for r in s5 if r[0] == r[3]]
    out["sheet5_self_pairs"] = {"connections": len(selfp), "synapses": sum(int(r[4]) for r in selfp)}
    out["sheet5_self_pairs_by_predicted_polarity"] = _by_category(selfp, 16, 4)
    srcs, tgts = {r[0] for r in s5}, {r[3] for r in s5}
    out["sheet5_distinct_sources"] = len(srcs)
    out["sheet5_distinct_targets"] = len(tgts)
    out["sheet5_distinct_names"] = len(srcs | tgts)
    out["sheet5_names_not_in_sheet1"] = sorted((srcs | tgts) - set(s1_names))
    out["sheet1_names_not_in_sheet5"] = sorted(set(s1_names) - (srcs | tgts))
    out["sheet5_rows_source_without_dominant_nt"] = sum(1 for r in s5 if r[1] in (0, None))

    # Polarity as the file encodes it: column Q "Predicted polarity".
    cats = _by_category(s5, 16, 4)
    out["sheet5_by_predicted_polarity_colQ"] = cats
    # Column S "Predicted synapse polarity": +E for '+', -E for '-', 0 for 'no pred', text 'complex'.
    s_num = [r[18] for r in s5 if isinstance(r[18], (int, float)) and not isinstance(r[18], bool)]
    out["sheet5_colS_value_kinds"] = dict(sorted(Counter(
        "positive" if (isinstance(v, (int, float)) and not isinstance(v, bool) and v > 0)
        else "negative" if (isinstance(v, (int, float)) and not isinstance(v, bool) and v < 0)
        else "zero" if v == 0 and not isinstance(v, bool) else f"text:{v}"
        for v in (r[18] for r in s5)
    ).items()))
    out["sheet5_colS_sum_positive"] = sum(v for v in s_num if v > 0)
    out["sheet5_colS_sum_negative"] = sum(v for v in s_num if v < 0)
    out["sheet5_colU_nopred_synapses_sum"] = sum(
        r[20] for r in s5 if isinstance(r[20], (int, float)) and not isinstance(r[20], bool)
    )
    out["sheet5_colT_unknown_synapses_sum"] = sum(
        r[19] for r in s5 if isinstance(r[19], (int, float)) and not isinstance(r[19], bool)
    )
    out["sheet5_neurotransmitter_complexity_colO_true"] = sum(1 for r in s5 if r[14] is True)
    out["sheet5_receptorial_complexity_colP_true"] = sum(1 for r in s5 if r[15] is True)

    # E:I side by side. The file's own figure is Summary!B10 (= B7/B8, synapse-weighted).
    e_syn, i_syn = cats["+"]["synapses"], cats["-"]["synapses"]
    e_con, i_con = cats["+"]["connections"], cats["-"]["connections"]
    out["ei_ratio_file_summary_B10"] = out["xlsx_summary_sheet_cells"]["B10"][0]
    out["ei_ratio_synapses_recomputed_from_colQ"] = _ratio(e_syn, i_syn)
    out["ei_ratio_connections_recomputed_from_colQ"] = _ratio(e_con, i_con)

    # Internal consistency: cached values vs the workbook's own formula logic (via sheets 1 and 4).
    nt = {r[0]: (r[1] if r[1] is not None else 0, r[2] if r[2] is not None else 0) for r in s1}
    rt = {r[0]: r[1:7] for r in s4}
    bad = Counter()
    for r in s5:
        if (r[1], r[2]) != nt[r[0]]:
            bad["B:C"] += 1
        g = []
        for k, name in enumerate(("Glu", "Glu", "ACh", "ACh", "GABA", "GABA")):
            hit = any(isinstance(x, str) and name.lower() in x.lower() for x in (r[1], r[2]))
            g.append(bool(hit and rt[r[3]][k] > 0))
        if tuple(g) != tuple(r[6:12]):
            bad["G:L"] += 1
        if _nt_r_category(g) != r[16]:
            bad["Q"] += 1
    out["sheet5_rows_inconsistent_with_own_formulas"] = dict(sorted(bad.items()))

    # How many data rows fall outside Excel's stale named ranges.
    m = re.search(r"\$(\d+)$", out["xlsx_defined_names"].get("herm_full_edgelist@5", ""))
    if m:
        out["sheet5_data_rows_after_named_range_herm_full_edgelist_end"] = sum(1 for i in data_idx if i > int(m.group(1)))

    # Any literal '4:1' anywhere in the workbook's text cells.
    hits = 0
    for name in vwb.sheetnames:
        for row in vwb[name].iter_rows(values_only=True):
            hits += sum(1 for v in row if isinstance(v, str) and "4:1" in v)
    out["xlsx_text_cells_containing_4to1"] = hits
    vwb.close()
    fwb.close()
    return s5, set(s1_names)


def _read_semicolon_csv(path):
    raw = path.read_bytes()
    text = raw.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=";"))
    fmt = {
        "bytes": len(raw),
        "utf8_bom": raw.startswith(b"\xef\xbb\xbf"),
        "crlf_line_breaks": raw.count(b"\r\n"),
        "bare_lf_line_breaks": raw.count(b"\n") - raw.count(b"\r\n"),
        "final_newline": raw.endswith(b"\n"),
    }
    return rows, fmt


def _xlsx_row_as_text(r):
    """Sheet-5 row A..Q in the CSVs' text spelling: booleans as TRUE/FALSE, numbers as digits."""
    return tuple("TRUE" if v is True else "FALSE" if v is False else str(v) for v in r[:17])


def _csv_counts(out: dict, raw, s5) -> None:
    xlsx_rows = [_xlsx_row_as_text(r) for r in s5]
    xlsx_by_pair = {(r[0], r[3]): r for r in xlsx_rows}

    # Sept 2020 revision.
    rows, fmt = _read_semicolon_csv(raw / REV_CSV)
    hdr, data = rows[0], rows[1:]
    out["rev2020_csv_format"] = fmt
    out["rev2020_csv_header"] = hdr
    out["rev2020_csv_data_rows"] = len(data)
    out["rev2020_csv_rows_by_type"] = dict(sorted(Counter(r[5] for r in data).items()))
    out["rev2020_csv_synapse_sum"] = sum(int(r[4]) for r in data)
    out["rev2020_csv_by_pred_pol"] = _by_category(data, 16, 4)
    out["rev2020_csv_equals_xlsx_sheet5_cols_A_to_Q_in_order"] = [tuple(r) for r in data] == xlsx_rows

    # Jan 2020 version.
    rows, fmt = _read_semicolon_csv(raw / JAN_CSV)
    hdr, data = rows[0], rows[1:]
    out["jan2020_csv_format"] = fmt
    out["jan2020_csv_header"] = hdr
    out["jan2020_csv_data_rows"] = len(data)
    out["jan2020_csv_rows_by_type"] = dict(sorted(Counter(r[5] for r in data).items()))
    out["jan2020_csv_rows_with_empty_type"] = [r for r in data if r[5] == ""]
    out["jan2020_csv_boolean_tokens_cols_G_to_L"] = dict(sorted(Counter(v for r in data for v in r[6:12]).items()))
    out["jan2020_csv_boolean_tokens_cols_G_to_L_by_type"] = {
        t: dict(sorted(Counter(v for r in data if r[5] == t for v in r[6:12]).items()))
        for t in sorted({r[5] for r in data})
    }
    out["jan2020_csv_rows_with_any_IGAZ_in_cols_G_to_L_by_type"] = {
        t: sum(1 for r in data if r[5] == t and "IGAZ" in r[6:12]) for t in sorted({r[5] for r in data})
    }
    chem = [r for r in data if r[5] == "chemical"]
    elec = [r for r in data if r[5] == "electrical"]
    out["jan2020_csv_electrical_rows_IGAZ_by_column_G_to_P"] = {
        hdr[j]: sum(1 for r in elec if r[j] == "IGAZ") for j in range(6, 16)
    }
    out["jan2020_csv_chemical_synapse_sum"] = sum(int(r[4]) for r in chem)
    out["jan2020_csv_chemical_by_pred_pol"] = _by_category(chem, 16, 4)
    out["jan2020_csv_ei_ratio_synapses"] = _ratio(
        out["jan2020_csv_chemical_by_pred_pol"]["+"]["synapses"], out["jan2020_csv_chemical_by_pred_pol"]["-"]["synapses"]
    )
    out["jan2020_csv_chemical_duplicate_rows"] = len(chem) - len({(r[0], r[3]) for r in chem})
    # Electrical rows, counted as rows (a duplicated (source, target) would count twice; see *_duplicate_rows).
    out["jan2020_csv_electrical_rows"] = len(elec)
    out["jan2020_csv_electrical_weight_sum"] = sum(int(r[4]) for r in elec)
    e_list = [(r[0], r[3], int(r[4])) for r in elec]
    e_pairs = Counter((a, b) for a, b, _ in e_list)
    e_weight = {}
    for a, b, w in e_list:
        e_weight.setdefault((a, b), []).append(w)
    out["jan2020_csv_electrical_duplicate_rows"] = sum(n - 1 for n in e_pairs.values())
    out["jan2020_csv_electrical_self_rows"] = sum(1 for a, b, _ in e_list if a == b)
    out["jan2020_csv_electrical_unordered_pairs"] = len({frozenset((a, b)) for a, b, _ in e_list})
    out["jan2020_csv_electrical_rows_whose_reverse_is_listed"] = sum(
        1 for a, b, _ in e_list if a != b and (b, a) in e_pairs
    )
    out["jan2020_csv_electrical_rows_listed_in_one_direction_only"] = sorted(
        [a, b, w] for a, b, w in e_list if a != b and (b, a) not in e_pairs
    )
    asym = sorted(
        [a, b, w, sum(e_weight[(b, a)])] for a, b, w in e_list
        if a != b and (b, a) in e_pairs and w != sum(e_weight[(b, a)])
    )
    out["jan2020_csv_electrical_rows_with_reverse_of_different_weight"] = len(asym)
    out["jan2020_csv_electrical_unordered_pairs_with_different_weights_per_direction"] = len(
        {frozenset((a, b)) for a, b, _, _ in asym}
    )
    out["jan2020_csv_electrical_rows_with_reverse_of_different_weight_list"] = asym
    out["jan2020_csv_electrical_pred_pol"] = dict(sorted(Counter(r[16] for r in elec).items()))
    n_chem = {r[0] for r in chem} | {r[3] for r in chem}
    n_elec = {r[0] for r in elec} | {r[3] for r in elec}
    out["jan2020_csv_distinct_names_chemical"] = len(n_chem)
    out["jan2020_csv_distinct_names_electrical"] = len(n_elec)
    out["jan2020_csv_distinct_names_chemical_or_electrical"] = len(n_chem | n_elec)
    out["jan2020_csv_names_electrical_only"] = sorted(n_elec - n_chem)

    # Jan vs xlsx, chemical rows. Hungarian booleans are mapped to TRUE/FALSE for the comparison only.
    tr = {"IGAZ": "TRUE", "HAMIS": "FALSE"}
    jan = [tuple(tr.get(v, v) for v in r) for r in chem]
    out["jan2020_vs_xlsx_same_pairs_and_order"] = [(r[0], r[3]) for r in jan] == [(r[0], r[3]) for r in xlsx_rows]
    out["jan2020_vs_xlsx_weight_differences"] = sum(
        1 for r in jan if (r[0], r[3]) in xlsx_by_pair and r[4] != xlsx_by_pair[(r[0], r[3])][4]
    )
    diff_cols = Counter()
    changed = []
    for r in jan:
        x = xlsx_by_pair[(r[0], r[3])]
        for j, (u, v) in enumerate(zip(r, x)):
            if u != v:
                diff_cols[hdr[j]] += 1
        if r[16] != x[16]:
            changed.append([r[0], r[3], int(r[4]), r[16], x[16]])
    out["jan2020_vs_xlsx_cells_differing_by_column"] = dict(sorted(diff_cols.items()))
    out["jan2020_vs_xlsx_pred_pol_changes"] = sorted(changed)
    out["jan2020_vs_xlsx_nt_complex_spelling"] = dict(sorted(Counter(
        f"{r[14]}->{xlsx_by_pair[(r[0], r[3])][14]}" for r in jan if r[14] != xlsx_by_pair[(r[0], r[3])][14]
    ).items()))


def _read_edgelist(path):
    """(header, rows as stripped [source, target, weight, type], raw-format facts) of a comma edge list.

    Whitespace is stripped only to read the fields; no other change is made to names.
    """
    raw_bytes = path.read_bytes()
    rows = list(csv.reader(io.StringIO(raw_bytes.decode("utf-8"), newline="")))
    hdr, data = rows[0], rows[1:]
    fmt = {
        "bytes": len(raw_bytes),
        "crlf_line_breaks": raw_bytes.count(b"\r\n"),
        "bare_lf_line_breaks": raw_bytes.count(b"\n") - raw_bytes.count(b"\r\n"),
        "final_newline": raw_bytes.endswith(b"\n"),
        "rows_with_space_padded_names": sum(1 for r in data if r[0] != r[0].strip() or r[1] != r[1].strip()),
    }
    return hdr, [[x.strip() for x in r] for r in data], fmt


def _triples(rows, keep):
    """Counter of (source, target, weight) over rows for which keep(row) is true (a multiset: duplicates count)."""
    return Counter((r[0], r[1], int(r[2])) for r in rows if keep(r))


def _layer(rows, typ, s1_names):
    """Neuron-to-neuron rows of one Type, both endpoints in the xlsx's sheet-1 list, as a multiset."""
    return _triples(rows, lambda r: r[3] == typ and r[0] in s1_names and r[1] in s1_names)


def _diff_rows(a: Counter, b: Counter):
    """Rows in multiset a but not in b, as sorted lists."""
    return sorted(list(t) for t in (a - b).elements())


def _edgelist_basic(out: dict, prefix: str, hdr, d, fmt) -> None:
    out[f"{prefix}_header"] = hdr
    out[f"{prefix}_data_rows"] = len(d)
    out[f"{prefix}_format"] = fmt
    out[f"{prefix}_rows_by_type"] = dict(sorted(Counter(r[3] for r in d).items()))
    out[f"{prefix}_weight_sum_by_type"] = {t: sum(int(r[2]) for r in d if r[3] == t) for t in sorted({r[3] for r in d})}
    out[f"{prefix}_duplicate_rows"] = len(d) - len({(r[0], r[1], r[3]) for r in d})


def _edgelist_counts(out: dict, raw, s5, s1_names: set) -> None:
    sheet5 = Counter((r[0], r[3], int(r[4])) for r in s5)
    jan_rows, _ = _read_semicolon_csv(raw / JAN_CSV)
    jan = [[r[0], r[3], r[4], r[5]] for r in jan_rows[1:] if r[5] in ("chemical", "electrical")]
    jan_chem = _triples(jan, lambda r: r[3] == "chemical")
    jan_elec = _triples(jan, lambda r: r[3] == "electrical")

    # The acquired ConnectomeToolbox copy (blob 89ffc2a2).
    hdr, d, fmt = _read_edgelist(raw / EDGELIST)
    _edgelist_basic(out, "edgelist", hdr, d, fmt)
    out["edgelist_rows_with_space_padded_names"] = fmt["rows_with_space_padded_names"]
    names = {r[0] for r in d} | {r[1] for r in d}
    out["edgelist_distinct_names"] = len(names)
    out["edgelist_names_not_in_xlsx_sheet1"] = len(names - s1_names)
    out["edgelist_names_not_in_xlsx_sheet1_with_uppercase_initial"] = sorted(
        n for n in names - s1_names if n[:1].isupper()
    )
    rule_a = _layer(d, "chemical", s1_names)
    rule_b = _triples(d, lambda r: r[3] == "chemical" and r[0][:1].isupper() and r[1][:1].isupper())
    out["edgelist_chemical_both_ends_in_xlsx_sheet1"] = {
        "connections": sum(rule_a.values()), "synapses": sum(w * n for (_, _, w), n in rule_a.items()),
        "equals_sheet5": rule_a == sheet5,
    }
    out["edgelist_chemical_both_ends_uppercase_initial"] = {
        "connections": sum(rule_b.values()), "synapses": sum(w * n for (_, _, w), n in rule_b.items()),
        "equals_sheet5": rule_b == sheet5,
    }
    elec = _layer(d, "electrical", s1_names)
    out["edgelist_electrical_both_ends_in_xlsx_sheet1"] = {
        "rows": sum(elec.values()), "weight": sum(w * n for (_, _, w), n in elec.items()),
    }
    out["edgelist_electrical_rows_missing_from_jan2020_csv"] = _diff_rows(elec, jan_elec)
    out["jan2020_csv_electrical_rows_missing_from_edgelist"] = _diff_rows(jan_elec, elec)

    # The pre-edit OpenWorm copy (CElegansNeuroML 3808b6b, blob a1313b0e).
    hdr_p, p, fmt_p = _read_edgelist(raw / EDGELIST_PRE_EDIT)
    _edgelist_basic(out, "edgelist_pre_edit", hdr_p, p, fmt_p)
    full = Counter((r[0], r[1], int(r[2]), r[3]) for r in d)
    full_p = Counter((r[0], r[1], int(r[2]), r[3]) for r in p)
    out["edgelist_rows_not_in_edgelist_pre_edit"] = _diff_rows(full, full_p)
    out["edgelist_pre_edit_rows_not_in_edgelist"] = _diff_rows(full_p, full)
    p_chem, p_elec = _layer(p, "chemical", s1_names), _layer(p, "electrical", s1_names)
    out["edgelist_pre_edit_chemical_both_ends_in_xlsx_sheet1"] = {
        "rows": sum(p_chem.values()), "weight": sum(w * n for (_, _, w), n in p_chem.items()),
        "equals_jan2020_csv_chemical": p_chem == jan_chem, "equals_sheet5": p_chem == sheet5,
    }
    out["edgelist_pre_edit_electrical_both_ends_in_xlsx_sheet1"] = {
        "rows": sum(p_elec.values()), "weight": sum(w * n for (_, _, w), n in p_elec.items()),
        "equals_jan2020_csv_electrical": p_elec == jan_elec,
    }
    out["edgelist_pre_edit_neuron_rows_in_same_order_as_jan2020_csv"] = [
        (r[0], r[1], int(r[2]), r[3]) for r in p if r[0] in s1_names and r[1] in s1_names
    ] == [(r[0], r[1], int(r[2]), r[3]) for r in jan]

    # The toolbox's MODIFIED copy (blob 5065baf7): evidence that it is not the Fenyves input.
    hdr_m, m, fmt_m = _read_edgelist(raw / EDGELIST_MODIFIED)
    _edgelist_basic(out, "edgelist_modified", hdr_m, m, fmt_m)
    full_m = Counter((r[0], r[1], int(r[2]), r[3]) for r in m)
    out["edgelist_modified_rows_not_in_edgelist"] = _diff_rows(full_m, full)
    out["edgelist_rows_not_in_edgelist_modified"] = _diff_rows(full, full_m)
    m_chem = _layer(m, "chemical", s1_names)
    m_pairs = {(a, b): w for (a, b, w) in m_chem}
    s5_pairs = {(a, b): w for (a, b, w) in sheet5}
    out["edgelist_modified_chemical_both_ends_in_xlsx_sheet1"] = {
        "rows": sum(m_chem.values()), "weight": sum(w * n for (_, _, w), n in m_chem.items()),
        "equals_sheet5": m_chem == sheet5,
    }
    out["edgelist_modified_chemical_pairs_not_in_sheet5"] = sorted([a, b, w] for (a, b), w in m_pairs.items() if (a, b) not in s5_pairs)
    out["sheet5_pairs_not_in_edgelist_modified"] = sorted([a, b, w] for (a, b), w in s5_pairs.items() if (a, b) not in m_pairs)
    out["edgelist_modified_vs_sheet5_weight_differences"] = sorted(
        [a, b, s5_pairs[(a, b)], w] for (a, b), w in m_pairs.items() if (a, b) in s5_pairs and s5_pairs[(a, b)] != w
    )


def record_counts() -> dict:
    """Counts computed only from data/raw/fenyves2020/. No renaming, no biological filtering.

    Workbook values are Excel's cached results; formulas are reported next to
    them where they exist. Definitions:

    - ``xlsx_summary_sheet_cells`` / ``xlsx_user_guide_cells`` /
      ``xlsx_sheet4_results_block_cells``: every non-empty cell, as
      [cached value, formula or null], keyed by cell reference.
    - ``sheet5_data_rows``: rows >= 3 of "5. Sign prediction" with a non-empty
      column A. One row = one directed connection Source (A) -> Target (D).
    - ``sheet5_edge_weight_sum``: sum of column E ("Edge weight", synapses).
    - ``sheet5_by_predicted_polarity_colQ``: rows ("connections") and sum of
      E ("synapses") per value of column Q ("Predicted polarity"), the file's
      own four categories '+', '-', 'complex', 'no pred'.
    - ``sheet5_colS_*`` / ``colT`` / ``colU``: the file's own signed-synapse
      columns, summed as encoded (S: +E, -E, 0 or 'complex'; T: -E for
      complex; U: E for no pred).
    - E:I, side by side: ``ei_ratio_file_summary_B10`` is the file's figure
      (Summary!B10 = B7/B8, synapse-weighted). ``..._synapses_recomputed`` is
      the same quotient from column Q; ``..._connections_recomputed`` counts
      rows instead of synapses. Self-pairs are included in all three; their
      per-category share is ``sheet5_self_pairs_by_predicted_polarity``.
    - ``sheet5_first_data_row_formulas``: the formula text of every formula
      cell in the first data row of sheet 5 (each column's shared-formula
      anchor), as written in the file.
    - ``sheet5_rows_inconsistent_with_own_formulas``: rows whose cached B:C,
      G:L or Q differ from the workbook's own formulas re-evaluated on
      sheets 1 and 4 (Excel SEARCH = case-insensitive substring).
    - ``sheet4_cells_hardcoded_not_formula_by_column``: receptor-count cells
      in sheet 4 that hold a typed value instead of a formula. Array (CSE)
      formulas count as formulas; ``sheet4_array_formula_cells_by_column``
      counts them separately.
    - ``sheet4_colI_first_row_formula``: the formula text of sheet 4 column I
      ("multi-receptor complexity") in the first neuron row, as openpyxl reads
      it; ``sheet4_colI_rows_not_matching_first_row_formula_pattern`` counts
      neuron rows whose column-I formula is not that formula with the first
      row's number replaced by the row's own.
    - ``sheet4_cols_B_to_G_cached_values_differing_from_sheets2_3_reevaluation``:
      cells of B..G whose cached value differs from the cell's own formula
      re-evaluated on sheets 2 and 3 (the class ranges are read from the
      workbook's defined names; COUNTIF as case-insensitive equality).
    - ``rev2020_*``: ``Data (rev 12-09-2020).csv``; ``jan2020_*``: ``data.csv``.
      "equals xlsx" compares text with booleans written TRUE/FALSE; the Jan
      file's IGAZ/HAMIS are mapped to TRUE/FALSE only for that comparison.
      Electrical statistics are computed over the row list, so a duplicated
      (source, target) row would be counted (see ``*_duplicate_rows``).
    - ``edgelist_*``: ``herm_full_edgelist.csv`` (toolbox, blob 89ffc2a2);
      ``edgelist_pre_edit_*``: the CElegansNeuroML 3808b6b copy (blob
      a1313b0e); ``edgelist_modified_*``: ``herm_full_edgelist_MODIFIED.csv``.
      Fields are whitespace-stripped to be read at all. Layers are compared
      as multisets of (source, target, weight) rows. Two endpoint rules are
      reported side by side for deriving sheet 5: both names in the xlsx's
      sheet-1 list (302 names), or both names starting with an upper-case
      letter.
    """
    raw = RAW_ROOT / SOURCE_ID
    out: dict = {}
    with warnings.catch_warnings():
        # Read-only openpyxl parses sheets lazily and warns about an Excel
        # extension it does not support; cell values are unaffected.
        warnings.filterwarnings("ignore", message="Unknown extension", category=UserWarning)
        s5, s1_names = _xlsx_counts(out, raw)
    _csv_counts(out, raw, s5)
    _edgelist_counts(out, raw, s5, s1_names)
    return out
