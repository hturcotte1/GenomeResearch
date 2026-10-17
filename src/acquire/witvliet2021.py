"""Witvliet et al. 2021, Nature 596:257-261: brain connectomes of 8 isogenic animals, L1 to adult.

The two adult datasets (7 and 8), the ones Randi et al. 2023 used, are
acquired in every version found. Datasets 1-6 (L1-L3) are acquired only in
the pre-revision NemaNode version, and only for per-dataset counts (how many
classified connections each dataset contains; dataset-1 connections followed
into the adults); their other copies are listed in
docs/provenance/witvliet2021.md as available but not acquired. The primary hosts (nature.com, nemanode.org,
wormwiring.org, zenodo.org) are blocked in this environment; see that document.

Four secondary distributions are acquired, because they hold three different
versions of the same two datasets plus the paper's own tables:

- ``aconnectome_witvliet_2020_{7,8}.csv``: members of the ``wormneuroatlas``
  0.0.7.3 wheel. Tab-separated despite the ``.csv`` suffix, no trailing
  newline, header ``pre post type synapses``. ``wormneuroatlas`` reads them
  (``NeuroAtlas.aconn_sources``) only when called with non-default
  parameters; by default it loads ``aconnectome_default.h5``, which was
  compiled from them (reproduced in Phase 1). Randi et al.'s analysis
  repository ``leiferlab/pumpprobe`` @ ``1dbc5e0a`` carries byte-identical
  copies (same git blob ids) and its ``Funatlas`` reads them directly.
- ``witvliet_2020_{7,8} adult.xlsx``: from the OpenWorm ConnectomeToolbox at
  the pinned commit. The toolbox says they are the WormWiring spreadsheets
  (``cect/validation/WitvlietEtAl2021.md``). One sheet, same four columns.
  The same commit's ``WitvlietDataReader{7,8}_expected_data.yaml``,
  ``WitvlietEtAl2021.md`` and ``docs/Witvliet_2021.md`` are acquired as
  evidence: the toolbox's expected counts, its statement of what it read from
  the Nature supplement, and its paraphrase of the paper text.
- ``nemanode_<commit>__*.json``: the source data of nemanode.org, from the Zhen
  lab's GitHub repository ``zhenlab-ltri/NemaNode`` (all data commits by
  D. Witvliet). ``git log --all --follow`` gives this history for datasets 7
  and 8: ``8a50236`` (2020-05-14, root commit, as ``TEM_adult.head.json`` /
  ``SEM_adult.head.json``); ``b6dfd28`` (2020-05-18, "Fixed bug in data where
  two synapses were counted only once", new content); ``a0ff6f3``
  (2020-05-18, directory move); ``859641e`` (2020-05-22, "Renamed datasets to
  be consistent with manuscript", a pure rename to ``witvliet_2020_{7,8}.json``,
  no content change); ``baf8ec8`` (2021-03-15, "Updated witvliet datasets after
  paper revision", new content). Pinned: ``8a50236`` (the first version),
  ``859641e`` (the second, "pre-revision", content as of ``b6dfd28``; also
  datasets 1-6 and ``head.annotations.json`` from this commit) and ``baf8ec8``
  (the published version). The blobs at ``baf8ec8`` are unchanged at the
  repository's master head ``6b7ca03`` (2022-03-16). ``neurons.json`` is
  NemaNode's cell table (name, class, type code, transmitter code); its blob
  is identical at every commit named here. ``head.annotations.json`` (one
  blob from ``8a50236`` to ``6b7ca03``) is NemaNode's per-connection
  classification (stable / variable / increase / decrease / postembryonic).
- ``synapse_count_matrices.xlsx``, ``synapse_list.csv``,
  ``connection_classifications.csv``, ``_neurons.csv``: the ``tables/``
  folder of the first author's figure-generation repository
  ``dwitvliet/nature2021`` at its master head ``0646af9`` (2021-06-01, "Fixed
  stats for publication"). The commits that added them are titled "Added
  export functions for supplementary tables", "Added supplementary tables" and
  "Added missing supplementary info" (2021-02-10/11/12). Whether they are
  byte-identical to the Nature supplementary files is unverified. They cover
  all 8 datasets (one file each); only datasets 7 and 8 are counted in detail.
  The same repository's ``data/nemanode/witvliet_2020_{7,8}.json`` have the
  same git blob ids as NemaNode ``baf8ec8``, i.e. the same bytes.

Facts checked by ``record_counts`` that later phases must respect:

- The wheel CSVs equal the PRE-revision NemaNode data (859641e), edge for edge
  and weight for weight. The toolbox xlsx equal the POST-revision data
  (baf8ec8), apart from gap-junction entries to unidentified fragments, which
  the xlsx omit. Randi et al.'s code (wormneuroatlas, pumpprobe) therefore
  uses the pre-revision version.
- The root-commit version (8a50236) is a third version, equal to neither.
- The paper's own ``synapse_count_matrices.xlsx`` (datasets 7, 8) and
  ``synapse_list.csv`` agree with the post-revision data, not with the wheel
  CSVs. Matrices are laid out [post row, pre column]. Their cell set has no
  ``excgl`` row or column.
- NemaNode's ``head.annotations.json`` (pre-revision classification) gives
  3,113 classified (non-post-embryonic) pairs, and their presence in the
  859641e datasets 1-8 averages 1,292.0 per dataset; the authors' Feb-2021
  ``connection_classifications.csv`` gives 3,122 and 1,298.625.
- Chemical rows are oriented pre -> post (presynaptic -> postsynaptic).
  Electrical rows have no direction. NemaNode JSON lists each gap junction in
  both directions; the CSV and xlsx list each unordered pair once (the xlsx
  always with pre <= post as strings, the CSV in no fixed order).
- A connection's weight is the number of entries in its ``syn`` list
  (NemaNode's own rule, ``populate-connections.js``: ``synapseCount =
  syn.length``). That equals the CSV/xlsx ``synapses`` column. A polyadic
  synapse is counted once per postsynaptic partner, and a connector id can be
  listed two or three times within one connection, so the weight sum is not a
  count of distinct synapses. Both are reported.
- Names are Witvliet/NemaNode names: AWCL/AWCR (not AWCON/AWCOFF), muscles as
  ``BWM-DL01``... (zero-padded), glia and other cells (CEPsh*, GLR*, excgl),
  and in the post-revision JSON the non-cell names ``NR_fragment``,
  ``vncfrag``, ``Fragment``, ``VAn``.
"""

from __future__ import annotations

import json
from collections import Counter

from .common import RAW_ROOT, ArchiveMember, RemoteFile, acquire_source, verify_pypi_digest
from .distributions import (
    WORMNEUROATLAS_PACKAGE,
    WORMNEUROATLAS_VERSION,
    WORMNEUROATLAS_WHEEL,
    connectome_toolbox_file,
)

SOURCE_ID = "witvliet2021"
CITATION = (
    "Witvliet D, Mulcahy B, Mitchell JK, et al. Connectomes across development "
    "reveal principles of brain maturation. Nature 596, 257-261 (2021). "
    "doi:10.1038/s41586-021-03778-8"
)
PRIMARY_PAPER = "https://www.nature.com/articles/s41586-021-03778-8"
PRIMARY_PORTAL = "https://nemanode.org/ (dataset ids witvliet_2020_7, witvliet_2020_8)"
PRIMARY_WORMWIRING = "https://wormwiring.org/pages/witvliet.html"

# The Zhen lab's NemaNode repository: the source data loaded into nemanode.org.
# Declared here, not in distributions.py, because only this source uses it.
NEMANODE_REPO = "zhenlab-ltri/NemaNode"
NEMANODE_INITIAL = "8a50236db4c5888e7fe0c1d018e5eff0729e4c88"  # 2020-05-14, root commit
NEMANODE_PREREVISION = "859641e20748f5313a99eaa862227ae9f6d369ca"  # 2020-05-22, rename only; content of b6dfd28 (2020-05-18)
NEMANODE_REVISED = "baf8ec877455890d089d965d1a0b3ea0db110e9f"  # 2021-03-15, D. Witvliet
NEMANODE_MASTER_HEAD = "6b7ca03f6d68c4ea1af3c120d1a47681edc74a90"  # 2022-03-16; same blobs as REVISED
_NN_DIR = "src/server/populate-db/raw-data"
_NN_DIR_INITIAL = "src/server/db/populate-db/raw-data"  # the same folder before a0ff6f3 moved it
# 859641e renamed TEM_adult.head.json -> witvliet_2020_7.json and SEM_adult.head.json ->
# witvliet_2020_8.json (100% similarity), which is what ties the root-commit files to datasets 7 and 8.
_NN_INITIAL_NAMES = {7: "TEM_adult.head.json", 8: "SEM_adult.head.json"}
# Old names of datasets 1-6 before the same rename (git show -M 859641e).
_NN_OLD_NAMES_1_6 = {1: "SEM_L1_3", 2: "TEM_L1_5", 3: "SEM_L1_4", 4: "SEM_L1_2", 5: "SEM_L2_2", 6: "TEM_L3"}

# Where the primary copy of each NemaNode version lives. Only the published
# version is (presumably) still served by nemanode.org; that is unverified.
_PRIMARY_REVISED = f"{PRIMARY_PORTAL}; paper {PRIMARY_PAPER}"
_PRIMARY_PREREVISION = (
    "none current: NemaNode repository content from b6dfd28 (2020-05-18) until baf8ec8 (2021-03-15), "
    "when the first author replaced it (nemanode.org now presumably serves the replacement; unverified); "
    f"paper {PRIMARY_PAPER}"
)
_PRIMARY_INITIAL = (
    "none current: NemaNode repository root commit 8a50236 (2020-05-14), replaced four days later by b6dfd28 "
    f"('Fixed bug in data where two synapses were counted only once'); paper {PRIMARY_PAPER}"
)

# The first author's figure-generation repository for the Nature paper.
NATURE2021_REPO = "dwitvliet/nature2021"
NATURE2021_COMMIT = "0646af9d25896ae660f97d462eab2d67282f5625"  # 2021-06-01, master head


def _nature2021_file(path: str, *, sha256: str, size: int, git_blob_sha1: str, note: str) -> RemoteFile:
    return RemoteFile(
        url=f"https://raw.githubusercontent.com/{NATURE2021_REPO}/{NATURE2021_COMMIT}/{path}",
        filename=path.rsplit("/", 1)[-1],
        sha256=sha256,
        size=size,
        distributor=f"GitHub {NATURE2021_REPO} @ {NATURE2021_COMMIT} (D. Witvliet, first author; figure code for the paper)",
        origin="secondary",
        primary_reference=f"{PRIMARY_PAPER} (Supplementary Information; correspondence to a specific supplementary file unverified)",
        git_blob_sha1=git_blob_sha1,
        note=note,
    )


def _nemanode_file(commit: str, path: str, *, filename: str, sha256: str, size: int,
                   git_blob_sha1: str, primary_reference: str, note: str) -> RemoteFile:
    return RemoteFile(
        url=f"https://raw.githubusercontent.com/{NEMANODE_REPO}/{commit}/{path}",
        filename=filename,
        sha256=sha256,
        size=size,
        distributor=f"GitHub {NEMANODE_REPO} @ {commit} (Zhen lab; source data of nemanode.org)",
        origin="secondary",
        primary_reference=primary_reference,
        git_blob_sha1=git_blob_sha1,
        note=note,
    )


# Datasets 1-6 at 859641e, needed only for the per-dataset classification
# counts (the 1,292 figure). (sha256, size, git blob), from a verified download.
_NN_PREREVISION_1_6 = {
    1: ("130f5b6cfefa824958374ec9f8f964202faf56baa8634f6b79d58e49962b84d0", 227839, "7cdc94fd3fa57c8084b0df81344787198e3eeab4"),
    2: ("83ad7f2410eeec903a2454882b35fd5210e6dae48bc652902cf3dc32665389aa", 309957, "e5ea2884d2a03fe364b082e86e4d1f7cde7472a0"),
    3: ("c2295e45df692a6d85004ebb70632b25f80fc2be59c4bec7ea3492912005c7c3", 314899, "da46eb7d634a02fead622adb4284ac32a0b01215"),
    4: ("d5c50a6f4248842a3155ef0d12bbe3b24ca0444d25a847404d41cf12cf077b62", 416671, "fe65dd81a6ab47bbb63fa0d84d0f6560fc6b001a"),
    5: ("879ae9ae44db9804b869841fd067e0d3eb584bf55faa2dd8c78c952d59269355", 596438, "cda01e99b9ab3fa406eec704b26f9aeef9bf219b"),
    6: ("c31e90ab7d677b7a19fad18b57aaa81ff7b6d1b4b13a08c23b0393503ecc1423", 580169, "e1e460f380ccc9cb910dcd616053769062f3b7ff"),
}


FILES = (
    # --- What Randi et al. used: wormneuroatlas 0.0.7.3 wheel members -------
    ArchiveMember(
        archive=WORMNEUROATLAS_WHEEL,
        member="wormneuroatlas/data/aconnectome_witvliet_2020_7.csv",
        filename="aconnectome_witvliet_2020_7.csv",
        sha256="1370d2a7e26a17e0765edfa834c95ca59ab44d8587efb5fda330432233074165",
        size=55102,
        primary_reference=_PRIMARY_PREREVISION,
        note=(
            "Dataset 7 (adult). Tab-separated. Read by wormneuroatlas NeuroAtlas only for non-default "
            "parameters (default: aconnectome_default.h5, compiled from it); byte-identical to "
            "leiferlab/pumpprobe @1dbc5e0a pumpprobe/aconnectome_witvliet_2020_7.csv (blob f312c4db), which "
            "Randi et al.'s Funatlas reads. Content equals NemaNode 859641e, not the post-revision data."
        ),
    ),
    ArchiveMember(
        archive=WORMNEUROATLAS_WHEEL,
        member="wormneuroatlas/data/aconnectome_witvliet_2020_8.csv",
        filename="aconnectome_witvliet_2020_8.csv",
        sha256="530a6d6b1c0c621d70a019e478ac79c8107cde475b55e79f225582f77b342ce4",
        size=55153,
        primary_reference=_PRIMARY_PREREVISION,
        note=(
            "Dataset 8 (adult). Tab-separated. Read by wormneuroatlas NeuroAtlas only for non-default "
            "parameters (default: aconnectome_default.h5, compiled from it); byte-identical to "
            "leiferlab/pumpprobe @1dbc5e0a pumpprobe/aconnectome_witvliet_2020_8.csv (blob b16614e2), which "
            "Randi et al.'s Funatlas reads. Content equals NemaNode 859641e, not the post-revision data."
        ),
    ),
    # --- WormWiring spreadsheets, via the OpenWorm ConnectomeToolbox ---------
    connectome_toolbox_file(
        "cect/data/witvliet_2020_7%20adult.xlsx",
        filename="witvliet_2020_7 adult.xlsx",
        sha256="d0fd2dd32adbe5f04ce90a2e5b74c01c9a831876527b0a2b8bd5b94edc86b037",
        size=53165,
        git_blob_sha1="9cec72011ebfb1052631fbb7f86d6bf3190ed421",
        primary_reference=f"{PRIMARY_WORMWIRING} (per the toolbox's cect/validation/WitvlietEtAl2021.md); portal {PRIMARY_PORTAL}",
        note=(
            "Dataset 7 (adult). LibreOffice 6.0.7.3 xlsx, zip entries dated 2022-09-20. "
            "Content equals NemaNode post-revision commit baf8ec8 minus gap junctions to unidentified fragments."
        ),
    ),
    connectome_toolbox_file(
        "cect/data/witvliet_2020_8%20adult.xlsx",
        filename="witvliet_2020_8 adult.xlsx",
        sha256="fdead89606257c1b26e57069fbe1de14c7696633b75b59b79e74c0bcc3497e62",
        size=53143,
        git_blob_sha1="b03bdbeddab70a0d70f6e430e8cdc584c46725c1",
        primary_reference=f"{PRIMARY_WORMWIRING} (per the toolbox's cect/validation/WitvlietEtAl2021.md); portal {PRIMARY_PORTAL}",
        note=(
            "Dataset 8 (adult). LibreOffice 6.0.7.3 xlsx, zip entries dated 2022-09-20. "
            "Content equals NemaNode post-revision commit baf8ec8 minus gap junctions to unidentified fragments."
        ),
    ),
    # --- NemaNode source data, post-revision (2021-03-15) --------------------
    _nemanode_file(
        NEMANODE_REVISED,
        f"{_NN_DIR}/connections/witvliet_2020_7.json",
        filename="nemanode_baf8ec8__witvliet_2020_7.json",
        sha256="350b247214763fe6b5a89b76a95cc018f7551c0446ad6501bd9c4774fb05fbae",
        size=854336,
        git_blob_sha1="201d50cd208f265d7b2819a87145cb2bcf718230",
        primary_reference=_PRIMARY_REVISED,
        note=(
            "Dataset 7, post-revision ('Updated witvliet datasets after paper revision', D. Witvliet, "
            "2021-03-15). Same blob at master head 6b7ca03. Per-synapse connector ids included."
        ),
    ),
    _nemanode_file(
        NEMANODE_REVISED,
        f"{_NN_DIR}/connections/witvliet_2020_8.json",
        filename="nemanode_baf8ec8__witvliet_2020_8.json",
        sha256="465aeb3c1f42d84273023cef8e5f5d8ebbcf63248550b2219db68fccb9787033",
        size=887526,
        git_blob_sha1="135b1806872eea14a5f7bb667901285d269df720",
        primary_reference=_PRIMARY_REVISED,
        note=(
            "Dataset 8, post-revision ('Updated witvliet datasets after paper revision', D. Witvliet, "
            "2021-03-15). Same blob at master head 6b7ca03. Per-synapse connector ids included."
        ),
    ),
    _nemanode_file(
        NEMANODE_REVISED,
        f"{_NN_DIR}/neurons.json",
        filename="nemanode_baf8ec8__neurons.json",
        sha256="c780e0b7cab28b311a7e0916a7b8e37b9cf72de152712bd4a5eb29fd9d595492",
        size=58575,
        git_blob_sha1="81da4572b20674dcc852cfa9fd041e66a3c7a73c",
        primary_reference=_PRIMARY_REVISED,
        note=(
            "NemaNode cell table (447 cells: name, classes, typ, nt, emb, inhead, intail). Identical blob at "
            "8a50236, b6dfd28, a0ff6f3, 859641e, baf8ec8 and 6b7ca03. Its typ/nt codes are NemaNode "
            "annotations, not Witvliet data."
        ),
    ),
    # --- NemaNode source data, pre-revision (content of b6dfd28, 2020-05-18; pinned at 859641e) ---
    _nemanode_file(
        NEMANODE_PREREVISION,
        f"{_NN_DIR}/connections/witvliet_2020_7.json",
        filename="nemanode_859641e__witvliet_2020_7.json",
        sha256="caff48ff6cde87983d616a67d1d5d1967871548a965f86722ab2b666fe0315c3",
        size=897485,
        git_blob_sha1="8ec3eab48416e4ea390fbcdf6be2085c14333f8d",
        primary_reference=_PRIMARY_PREREVISION,
        note=(
            "Dataset 7, pre-revision. Same blob since b6dfd28 (2020-05-18, as TEM_adult.head.json); 859641e "
            "(2020-05-22) only renamed it. Replaced in the repository by baf8ec8 (2021-03-15). Its chemical "
            "edges and weights, and gap junctions as unordered pairs, equal the wheel CSV."
        ),
    ),
    _nemanode_file(
        NEMANODE_PREREVISION,
        f"{_NN_DIR}/connections/witvliet_2020_8.json",
        filename="nemanode_859641e__witvliet_2020_8.json",
        sha256="c13c42b5ba662e0bddcb1060693e0f75d579d85ee5c81ca539c620271819b07a",
        size=932228,
        git_blob_sha1="c156b36c7d9efb28d7cc048cbe173f6bc651ee39",
        primary_reference=_PRIMARY_PREREVISION,
        note=(
            "Dataset 8, pre-revision. Same blob since b6dfd28 (2020-05-18, as SEM_adult.head.json); 859641e "
            "(2020-05-22) only renamed it. Replaced in the repository by baf8ec8 (2021-03-15). Its chemical "
            "edges and weights, and gap junctions as unordered pairs, equal the wheel CSV."
        ),
    ),
    *(
        _nemanode_file(
            NEMANODE_PREREVISION,
            f"{_NN_DIR}/connections/witvliet_2020_{k}.json",
            filename=f"nemanode_859641e__witvliet_2020_{k}.json",
            sha256=sha,
            size=size,
            git_blob_sha1=blob,
            primary_reference=_PRIMARY_PREREVISION,
            note=(
                f"Dataset {k} (juvenile), pre-revision. Same blob since b6dfd28 (as {_NN_OLD_NAMES_1_6[k]}.head.json); "
                "859641e only renamed it. Acquired only to count classified connections per dataset "
                "(the 1,292 figure); not used for edges."
            ),
        )
        for k, (sha, size, blob) in sorted(_NN_PREREVISION_1_6.items())
    ),
    _nemanode_file(
        NEMANODE_PREREVISION,
        f"{_NN_DIR}/annotations/head.annotations.json",
        filename="nemanode_859641e__head.annotations.json",
        sha256="65f1e2b629f6966312861be1f0b1455e7a5948e650dd388b0c5f6a1ac9124fdb",
        size=156426,
        git_blob_sha1="dae8701a1ffc09152261ce840f842d9bfdd46ad5",
        primary_reference=(
            f"https://nemanode.org/ connection annotations (same blob at master head {NEMANODE_MASTER_HEAD[:7]}; "
            f"whether nemanode.org serves it is unverified); paper {PRIMARY_PAPER}"
        ),
        note=(
            "NemaNode's connection classification: {category: [[pre, post], ...]} with categories stable, "
            "variable, increase, decrease, postembryonic. One blob from the root commit 8a50236 (2020-05-14) "
            "to 6b7ca03; a0ff6f3 only moved its folder; the 2021 revision (baf8ec8) did not change it."
        ),
    ),
    # --- NemaNode source data, root commit (2020-05-14), datasets 7 and 8 ----
    *(
        _nemanode_file(
            NEMANODE_INITIAL,
            f"{_NN_DIR_INITIAL}/connections/{_NN_INITIAL_NAMES[k]}",
            filename=f"nemanode_8a50236__{_NN_INITIAL_NAMES[k]}",
            sha256=sha,
            size=size,
            git_blob_sha1=blob,
            primary_reference=_PRIMARY_INITIAL,
            note=(
                f"Dataset {k} (adult), first version ('Initial commit after removing unpublished datasets'). "
                f"Replaced by b6dfd28 (2020-05-18); renamed to witvliet_2020_{k}.json in 859641e. A third "
                "version, equal to neither the pre-revision nor the published data."
            ),
        )
        for k, sha, size, blob in (
            (7, "4b9b5392603c621f7e76bb69880d314f401baffe6ffe7417ddfc106792435ab7", 827927,
             "fbe13f1e916cc9842c2877d3a7f4717e8e758283"),
            (8, "4de49ad034465df9cd1ff2e9d93a0a985e06ff1ca637b62adbdd365565f5bad7", 859998,
             "cf39c0d4114bb0584b2bd51fdc1c496653c65837"),
        )
    ),
    # --- ConnectomeToolbox validation files (evidence about the Nature supplement) ---
    connectome_toolbox_file(
        "cect/validation/WitvlietDataReader7_expected_data.yaml",
        sha256="e68d150f07cf3c30164e78607a43768188e4154d471c64d70c456fd5a81c9a95",
        size=440,
        git_blob_sha1="879672e7f78b97dda2dc7b0fd0ad7d2d67821f50",
        primary_reference=(
            "none other: written by the ConnectomeToolbox itself; per its WitvlietEtAl2021.md the chemical "
            f"weights listed were read from the Nature supplement 41586_2021_3778_MOESM4_ESM.xlsx ({PRIMARY_PAPER})"
        ),
        note=(
            "Toolbox expected values for its dataset-7 reader: 3 sampled chemical and 3 electrical weights, and "
            "total_nonzero_conns per synapse class. Evidence only; not a Witvliet file."
        ),
    ),
    connectome_toolbox_file(
        "cect/validation/WitvlietDataReader8_expected_data.yaml",
        sha256="1e99bfe82b5f1b4b1bf5f2d40360198aeac6c34739720fbdd496b83c2ba89bda",
        size=435,
        git_blob_sha1="041a0e3ef3e7e783001c568afea41f1464be6f8a",
        primary_reference=(
            "none other: written by the ConnectomeToolbox itself; per its WitvlietEtAl2021.md the chemical "
            f"weights listed were read from the Nature supplement 41586_2021_3778_MOESM4_ESM.xlsx ({PRIMARY_PAPER})"
        ),
        note=(
            "Toolbox expected values for its dataset-8 reader: 3 sampled chemical and 3 electrical weights, and "
            "total_nonzero_conns per synapse class. Evidence only; not a Witvliet file."
        ),
    ),
    connectome_toolbox_file(
        "cect/validation/WitvlietEtAl2021.md",
        sha256="452c0977a771a3c08bb3e5ac0dbc7ae9e400cd46c359d0a9a455a3ec8ba73f6c",
        size=1035,
        git_blob_sha1="7e5d091debb2ef06cf7f17ed4b31db0dd7df6a01",
        primary_reference="none other: ConnectomeToolbox documentation text",
        note=(
            "The toolbox's note on its Witvliet readers: names the Nature MOESM4 file, says the xlsx came from "
            "WormWiring, and says which expected weights were read from which file."
        ),
    ),
    connectome_toolbox_file(
        "docs/Witvliet_2021.md",
        sha256="4fad03369d8613eb4cae9745d6db556d35e4bc31f40e0af9d3cbae39b23eebc4",
        size=3263,
        git_blob_sha1="0a8961df039d5adb0c5559b0ad62f871b255998b",
        primary_reference=f"none other: ConnectomeToolbox summary of the paper ({PRIMARY_PAPER}), whose text it paraphrases",
        note=(
            "The toolbox's paraphrase of the paper; the only source reached for the figures 3,113 and 1,292 "
            "(paper text unreachable). Evidence only."
        ),
    ),
    # --- The paper's own tables (first author's repository) -----------------
    _nature2021_file(
        "tables/synapse_count_matrices.xlsx",
        sha256="bd725d1f6121abeb8ec08d0b80edad94d9e0fea954a8704ee07c7f4be7cebfe6",
        size=764283,
        git_blob_sha1="e2c29655de6932b1d9e9d968b16f7ed318d3e7e3",
        note=(
            "Chemical synapse count matrices, one sheet per dataset 1-8, rows = post, columns = pre, "
            "brain cells only; no gap junctions. Written by export_matrices.py (commit 0fa137d, 2021-02-10)."
        ),
    ),
    _nature2021_file(
        "tables/synapse_list.csv",
        sha256="36c431f25e72cd82946239a2eb555c1ed4494a2036092105bd37eb161ecf9340",
        size=1287223,
        git_blob_sha1="80fc81032f996e75349397c2dec505cd462ca48c",
        note=(
            "One row per (synapse, postsynaptic partner), datasets 1-8: dataset, pre, post, syn_id, sections, "
            "size (-1 = missing), post_is_branch. Commit d526424, 2021-02-11."
        ),
    ),
    _nature2021_file(
        "tables/connection_classifications.csv",
        sha256="da94863507a64da8d85d75a280ceff538d30f98a746390a7d0090607b0870f08",
        size=92737,
        git_blob_sha1="960ce494eee0352320bde8a3bc82a59fede4d713",
        note="The paper's per-connection classification (stable / variable / developmentally dynamic / post-embryonic). Commit 8b2ce2e, 2021-02-12.",
    ),
    _nature2021_file(
        "tables/_neurons.csv",
        sha256="2189e2ee262a1d1a81f1139fcc95c4b886681cfdd706dc5f77bd8cda2a49491d",
        size=2223,
        git_blob_sha1="bef93bfd871caa8c3fbea910741d01bc9c960ac8",
        note="The paper's brain cell classes (class, members, type, integration); tables.py omits CAN. Commit 65282f6, 2021-02-11.",
    ),
)

DATASETS = (7, 8)
VERSIONS = ("wna_csv", "cect_xlsx", "nemanode_8a50236", "nemanode_859641e", "nemanode_baf8ec8")
_NEMANODE_TYP = {0: "chemical", 2: "electrical"}
_POSTEMB_TABLE = "post-embryonic brain integration"  # connection_classifications.csv
_POSTEMB_NEMANODE = "postembryonic"  # head.annotations.json category


def acquire() -> list[dict]:
    pypi_check = verify_pypi_digest(
        WORMNEUROATLAS_PACKAGE,
        WORMNEUROATLAS_VERSION,
        WORMNEUROATLAS_WHEEL.filename,
        WORMNEUROATLAS_WHEEL.sha256,
    )
    extra = {
        f.filename: [f"archive {WORMNEUROATLAS_WHEEL.filename}: {pypi_check}"]
        for f in FILES
        if isinstance(f, ArchiveMember)
    }
    return acquire_source(SOURCE_ID, FILES, extra)


# ---------------------------------------------------------------------------
# Readers. Each returns rows (pre, post, type, weight) exactly as stored, plus,
# for NemaNode JSON, the connector-id lists. No renaming, no filtering.
# ---------------------------------------------------------------------------

def _read_tsv(path) -> tuple[list[str], list[tuple[str, str, str, int]], dict]:
    raw = path.read_bytes()
    text = raw.decode("ascii")
    lines = text.split("\n")
    header = lines[0].split("\t")
    rows = []
    for line in lines[1:]:
        if line == "":
            continue
        pre, post, typ, n = line.split("\t")
        rows.append((pre, post, typ, int(n)))
    facts = {
        "bytes_contain_CR": b"\r" in raw,
        "ends_with_newline": raw.endswith(b"\n"),
        "delimiter_is_tab": all(len(line.split("\t")) == 4 for line in lines if line),
    }
    return header, rows, facts


def _read_xlsx(path) -> tuple[list[str], list[tuple[str, str, str, int]], dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        sheetnames = list(wb.sheetnames)
        values = [tuple(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()
    header = [str(v) for v in values[0]]
    cell_types = Counter(tuple(type(v).__name__ for v in r) for r in values[1:])
    for i, r in enumerate(values[1:], start=2):
        if len(r) != 4 or not (all(isinstance(v, str) for v in r[:3]) and type(r[3]) is int):
            raise ValueError(f"{path.name} row {i}: expected (str, str, str, int), got {r!r}")
    rows = [(a, b, c, d) for a, b, c, d in values[1:]]
    facts = {
        "sheetnames": sheetnames,
        "data_row_cell_types": {"|".join(k): v for k, v in sorted(cell_types.items())},
    }
    return header, rows, facts


def _read_nemanode(path) -> tuple[list[tuple[str, str, str, int]], list[dict], dict]:
    entries = json.loads(path.read_text(encoding="utf-8"))
    typ_values = Counter(e["typ"] for e in entries)
    rows = [(e["pre"], e["post"], _NEMANODE_TYP.get(e["typ"], f"typ={e['typ']}"), len(e["syn"])) for e in entries]
    facts = {
        "entry_key_sets": sorted({",".join(sorted(e.keys())) for e in entries}),
        "typ_values": {str(k): v for k, v in sorted(typ_values.items())},
        "syn_values": {str(k): v for k, v in sorted(Counter(s for e in entries for s in e["syn"]).items())},
        "all_list_lengths_equal": all(
            len(e["syn"]) == len(e["ids"]) == len(e["pre_tid"]) == len(e["post_tid"]) for e in entries
        ),
    }
    return rows, entries, facts


def _load_all(raw):
    out = {}
    for k in DATASETS:
        h, rows, facts = _read_tsv(raw / f"aconnectome_witvliet_2020_{k}.csv")
        out[("wna_csv", k)] = {"header": h, "rows": rows, "facts": facts, "entries": None}
        h, rows, facts = _read_xlsx(raw / f"witvliet_2020_{k} adult.xlsx")
        out[("cect_xlsx", k)] = {"header": h, "rows": rows, "facts": facts, "entries": None}
        for commit, fname in (
            ("8a50236", f"nemanode_8a50236__{_NN_INITIAL_NAMES[k]}"),
            ("859641e", f"nemanode_859641e__witvliet_2020_{k}.json"),
            ("baf8ec8", f"nemanode_baf8ec8__witvliet_2020_{k}.json"),
        ):
            rows, entries, facts = _read_nemanode(raw / fname)
            out[(f"nemanode_{commit}", k)] = {"header": None, "rows": rows, "facts": facts, "entries": entries}
    return out


def _nemanode_chem(path) -> dict:
    """{(pre, post): len(syn)} for the chemical entries of one NemaNode file; a repeated pair raises."""
    out = {}
    for e in json.loads(path.read_text(encoding="utf-8")):
        if e["typ"] != 0:
            continue
        if (e["pre"], e["post"]) in out:
            raise ValueError(f"{path.name}: chemical pair {e['pre']}->{e['post']} listed twice")
        out[(e["pre"], e["post"])] = len(e["syn"])
    return out


def _cell_groups(names, cells: dict) -> dict:
    """Partition names by NemaNode neurons.json ``typ``. No biological judgement is added."""
    groups = {"typ_b_muscle": [], "typ_u_unknown": [], "typ_neuron_codes": [], "absent_from_neurons_json": []}
    for n in sorted(names):
        if n not in cells:
            groups["absent_from_neurons_json"].append(n)
        elif cells[n]["typ"] == "b":
            groups["typ_b_muscle"].append(n)
        elif cells[n]["typ"] == "u":
            groups["typ_u_unknown"].append(n)
        else:
            groups["typ_neuron_codes"].append(n)
    return groups


def _edge_dicts(rows):
    """Chemical {(pre, post): weight}; gap {frozenset: set of stored weights}. A repeated chemical pair raises."""
    chem = {}
    gap = {}
    for pre, post, typ, w in rows:
        if typ == "chemical":
            if (pre, post) in chem:
                raise ValueError(f"chemical pair {pre}->{post} stored twice; not summed silently")
            chem[(pre, post)] = w
        elif typ == "electrical":
            gap.setdefault(frozenset((pre, post)), set()).add(w)
    return chem, gap


def _fmt_pair(p) -> str:
    a = sorted(p)
    return f"{a[0]}--{a[-1]}"


def _compare(a_rows, b_rows) -> dict:
    ca, ga = _edge_dicts(a_rows)
    cb, gb = _edge_dicts(b_rows)
    out = {
        "chem_edges_only_in_a": sorted(f"{p}->{q}:{ca[(p, q)]}" for p, q in ca if (p, q) not in cb),
        "chem_edges_only_in_b": sorted(f"{p}->{q}:{cb[(p, q)]}" for p, q in cb if (p, q) not in ca),
        "chem_weight_differs_a_vs_b": sorted(
            f"{p}->{q}:{ca[(p, q)]}vs{cb[(p, q)]}" for p, q in ca if (p, q) in cb and ca[(p, q)] != cb[(p, q)]
        ),
        "gap_pairs_only_in_a": sorted(f"{_fmt_pair(p)}:{sorted(ga[p])}" for p in ga if p not in gb),
        "gap_pairs_only_in_b": sorted(f"{_fmt_pair(p)}:{sorted(gb[p])}" for p in gb if p not in ga),
        "gap_weight_differs_a_vs_b": sorted(
            f"{_fmt_pair(p)}:{sorted(ga[p])}vs{sorted(gb[p])}" for p in ga if p in gb and ga[p] != gb[p]
        ),
    }
    out["chem_identical"] = not (out["chem_edges_only_in_a"] or out["chem_edges_only_in_b"] or out["chem_weight_differs_a_vs_b"])
    out["gap_unordered_identical"] = not (out["gap_pairs_only_in_a"] or out["gap_pairs_only_in_b"] or out["gap_weight_differs_a_vs_b"])
    out["chem_synapse_sum_a_minus_b"] = sum(ca.values()) - sum(cb.values())
    return out


def record_counts() -> dict:
    """Counts computed from the raw files only, with no renaming or filtering applied.

    Keys are ``<version>_d<dataset>_<quantity>``. Versions: ``wna_csv`` (wheel
    CSV, what Randi used), ``cect_xlsx`` (toolbox/WormWiring xlsx),
    ``nemanode_8a50236`` (root-commit JSON, ``TEM_adult``/``SEM_adult``),
    ``nemanode_859641e`` (pre-revision JSON), ``nemanode_baf8ec8``
    (post-revision JSON). Datasets: 7 and 8.

    Row and edge counts:

    - ``rows_<type>``: stored rows (JSON entries) by type. JSON ``typ`` 0 is
      "chemical", 2 is "electrical" (NemaNode docs); any other value would be
      reported as ``typ=<n>``.
    - ``duplicate_keys``: rows sharing (pre, post, type) with another row.
    - ``chem_edges``: chemical rows = ordered (pre, post) pairs.
    - ``chem_self_edges``: chemical rows with pre == post.
    - ``chem_synapse_sum``: sum of chemical weights. Weight is the
      ``synapses`` column (CSV/xlsx) or ``len(syn)`` (JSON; NemaNode's rule).
      A polyadic synapse adds 1 to each postsynaptic partner's edge.
    - JSON only: ``chem_unique_connector_ids`` = distinct connector ids over
      all chemical entries (one per presynaptic site); ``chem_sum_unique_ids_per_edge``
      = per-edge distinct ids, summed; ``chem_entries_with_repeated_id`` =
      entries listing an id more than once; ``chem_syn_value_sum`` = sum of
      the ``syn`` values themselves (differs from ``chem_synapse_sum`` where a
      ``syn`` value is 2); ``chem_entries_with_syn_value_gt1`` lists them.
      ``chem_connector_id_min``/``_max``: range of the ``ids`` values.
      ``chem_records_minus_unique_ids_per_edge``: records beyond the first
      for an id repeated within one entry (= ``chem_synapse_sum`` minus
      ``chem_sum_unique_ids_per_edge``). ``chem_id_multiplicity_within_entry``:
      how many (entry, id) occur 1, 2, 3... times. ``chem_ids_occurring_3plus_times_within_entry``
      lists those with 3 or more. ``chem_identical_records_within_entry``:
      (id, pre_tid, post_tid) triples stored more than once in one entry.
    - ``gap_rows``: electrical rows as stored. ``gap_unordered_pairs``:
      distinct {a, b}. ``gap_self_pairs``: a == b. ``gap_pairs_in_both_orientations``:
      unordered pairs (a != b) stored as both a->b and b->a.
      ``gap_pairs_one_orientation_only``: a != b stored one way only.
      ``gap_pairs_weight_asymmetric``: pairs whose two stored directions carry
      different weights. ``gap_weight_sum_unordered``: sum over unordered pairs
      of the pair's weight (max over stored directions; equal when symmetric).
      ``gap_rows_pre_lt_post``: rows with pre < post as Python strings.
      JSON only: ``gap_unique_connector_ids``.

    Cells (``cells_*``): the distinct names in pre or post, all row types,
    partitioned by NemaNode ``neurons.json`` ``typ`` only: ``b`` (muscle),
    ``u`` (unknown; NemaNode puts glia (CEPsh), GLR cells, CANR and excgl here), any
    neuron code (s/i/m/n combinations), or absent from neurons.json. This is
    NemaNode's labelling, reported as-is; which cells count as neurons is an
    owner decision (Phase 2). ``chem_edges_neuron_coded_both_ends`` and its
    synapse sum restrict to rows whose two ends both have a neuron code.

    ``cmp_d<k>_<a>__vs__<b>``: edge-level comparisons. Chemical edges are
    compared as ordered pairs with their weights; gap junctions as unordered
    pairs with the set of stored weights. Entries are "PRE->POST:w" or
    "A--B:[w]" strings, sorted. ``..._weight_sum_syn_...`` repeats the
    NemaNode pre- vs post-revision comparison with weight = sum of ``syn``.

    ``cells_by_paper_type``: each name is mapped to its class by NemaNode
    ``neurons.json`` ``classes`` and the class to its type by the paper's
    ``_neurons.csv``; names with no such type (absent from neurons.json, or
    class not listed, e.g. CAN, which the authors' tables.py skips) are listed
    in ``cells_without_paper_type_names``.

    Paper tables (``paper_*``): ``_neurons.csv`` class and member counts by
    type and integration; ``synapse_count_matrices.xlsx`` per adult sheet:
    pre columns, post rows (named rows only), nonzero cells, sum, the type
    labels in sheet order, post rows per type label and the names under
    "Other"; ``synapse_list.csv`` rows, distinct ``syn_id`` and distinct
    (pre, post) per dataset (all 8), rows with size == -1;
    ``connection_classifications.csv`` rows by value, duplicate (pre, post)
    pairs, rows not "post-embryonic brain integration" (and how many are
    self-pairs), their distinct pairs, and, per dataset, how many of those
    pairs occur in that dataset's synapse list (sum, number of datasets,
    mean). ``cmp_d<k>_paper_matrix__vs__<v>`` and
    ``cmp_d<k>_paper_synapse_list__vs__nemanode_baf8ec8`` compare chemical
    edges only (the paper tables hold no gap junctions); for the synapse list
    the weight is the number of rows per (pre, post), and ``syn_id_sets_equal``
    compares its syn_ids with NemaNode's chemical connector ids.
    ``cmp_d<k>_paper_matrix__vs__<v>_restricted_to_matrix_cells_*`` repeats
    the matrix comparison like for like: version <v>'s chemical edges are
    first restricted to the matrix's cell set (pre among its columns, post
    among its rows); ``..._b_edges_outside`` lists the edges dropped and
    ``..._b_chem_edges`` / ``..._b_chem_synapse_sum`` count what remains.

    NemaNode history (``nemanode_head_annotations_*``, from
    ``head.annotations.json``): entries per category, distinct pairs,
    duplicates, pairs in more than one category, self pairs, and
    ``pairs_not_postembryonic`` (distinct pairs in every category except
    "postembryonic"). Their presence per dataset is counted against the
    chemical (pre, post) pairs of NemaNode 859641e datasets 1-8
    (``..._present_in_nemanode_859641e_chem_*``) and of ``synapse_list.csv``
    (``..._present_in_paper_synapse_list_*``), with sum, number of datasets
    and mean. ``cmp_nemanode_head_annotations__vs__paper_classifications_*``
    lists the non-post-embryonic pairs found in only one of the two
    classifications. ``nemanode_859641e_by_dataset_chem_edges`` /
    ``_chem_synapse_sum``: datasets 1-8, weight = len(syn).
    ``d1_edges_in_adult__<source>_*``: dataset-1 chemical edges of
    ``nemanode_859641e`` or ``paper_synapse_list``, how many are present in
    adult dataset 7 or 8 of the same source, the adult synapse sum over them,
    and that sum divided by the present count (``mean_on_present``) or by all
    dataset-1 edges (``mean_on_all_d1_edges``).

    ConnectomeToolbox evidence (``toolbox_*``): the lines of its paper
    summary ``docs/Witvliet_2021.md`` that state 3,113 or 1,292; whether its
    ``WitvlietEtAl2021.md`` names MOESM4 and says the chemical weights were
    read from the supplement; per dataset, the yaml's reader name,
    ``total_nonzero_conns`` and sampled weights per synapse class, and for
    each version (and the paper matrix) the sampled chemical weights it does
    not reproduce (empty = all reproduced). ``cect_xlsx_d<k>_chem_nonzero_cells``
    and ``..._gap_nonzero_cells_symmetric`` (2 x non-self pairs + self pairs)
    are the counts a matrix reader of the xlsx would give, for comparison
    with ``total_nonzero_conns``.
    """
    raw = RAW_ROOT / SOURCE_ID
    cells = {c["name"]: c for c in json.loads((raw / "nemanode_baf8ec8__neurons.json").read_text(encoding="utf-8"))}
    data = _load_all(raw)
    paper_neurons = _read_csv_dicts(raw / "_neurons.csv")
    paper_type = {r["class"]: r["type"] for r in paper_neurons}
    out: dict = {
        "neurons_json_cells": len(cells),
        "neurons_json_typ_values": {k: v for k, v in sorted(Counter(c["typ"] for c in cells.values()).items())},
    }

    for (version, k), d in sorted(data.items()):
        key = f"{version}_d{k}"
        rows = d["rows"]
        if d["header"] is not None:
            out[f"{key}_header"] = d["header"]
        for fk, fv in sorted(d["facts"].items()):
            out[f"{key}_file_{fk}"] = fv
        types = Counter(r[2] for r in rows)
        out[f"{key}_rows_total"] = len(rows)
        for t, n in sorted(types.items()):
            out[f"{key}_rows_{t}"] = n
        keyc = Counter((r[0], r[1], r[2]) for r in rows)
        out[f"{key}_duplicate_keys"] = sum(v - 1 for v in keyc.values() if v > 1)
        out[f"{key}_weight_min"] = min(r[3] for r in rows)

        names = {r[0] for r in rows} | {r[1] for r in rows}
        groups = _cell_groups(names, cells)
        out[f"{key}_cells_total"] = len(names)
        for g, members in groups.items():
            out[f"{key}_cells_{g}"] = len(members)
        out[f"{key}_cells_typ_u_unknown_names"] = groups["typ_u_unknown"]
        out[f"{key}_cells_absent_from_neurons_json_names"] = groups["absent_from_neurons_json"]
        out[f"{key}_cells_AWC_names"] = sorted(n for n in names if n.startswith("AWC"))
        ptypes = Counter()
        no_paper_type = []
        for n in sorted(names):
            t = paper_type.get(cells[n]["classes"]) if n in cells else None
            if t is None:
                no_paper_type.append(n)
            else:
                ptypes[t] += 1
        out[f"{key}_cells_by_paper_type"] = dict(sorted(ptypes.items()))
        out[f"{key}_cells_without_paper_type_names"] = no_paper_type
        neuron_coded = set(groups["typ_neuron_codes"])

        chem = [r for r in rows if r[2] == "chemical"]
        out[f"{key}_chem_edges"] = len(chem)
        out[f"{key}_chem_self_edges"] = sum(r[0] == r[1] for r in chem)
        out[f"{key}_chem_self_edge_names"] = sorted(f"{r[0]}:{r[3]}" for r in chem if r[0] == r[1])
        out[f"{key}_chem_synapse_sum"] = sum(r[3] for r in chem)
        out[f"{key}_chem_weight_max"] = max(r[3] for r in chem)
        nn = [r for r in chem if r[0] in neuron_coded and r[1] in neuron_coded]
        out[f"{key}_chem_edges_neuron_coded_both_ends"] = len(nn)
        out[f"{key}_chem_synapse_sum_neuron_coded_both_ends"] = sum(r[3] for r in nn)
        out[f"{key}_chem_edges_neuron_coded_both_ends_non_self"] = sum(r[0] != r[1] for r in nn)

        gap = [r for r in rows if r[2] == "electrical"]
        directed = {}
        for pre, post, _, w in gap:
            if (pre, post) in directed:
                raise ValueError(f"{key}: electrical row {pre}->{post} stored twice; not overwritten silently")
            directed[(pre, post)] = w
        pairs = {frozenset((a, b)) for a, b in directed}
        both = [p for p in pairs if len(p) == 2 and all((a, b) in directed for a in p for b in p if a != b)]
        one = [p for p in pairs if len(p) == 2 and p not in set(both)]
        asym = [p for p in both if len({directed[(a, b)] for a in p for b in p if a != b}) > 1]

        def pair_weight(p):
            if len(p) == 1:
                (a,) = tuple(p)
                return directed[(a, a)]
            ws = [directed[(a, b)] for a in p for b in p if a != b and (a, b) in directed]
            return max(ws)

        out[f"{key}_gap_rows"] = len(gap)
        out[f"{key}_gap_unordered_pairs"] = len(pairs)
        out[f"{key}_gap_self_pairs"] = sum(len(p) == 1 for p in pairs)
        out[f"{key}_gap_pairs_in_both_orientations"] = len(both)
        out[f"{key}_gap_pairs_one_orientation_only"] = len(one)
        out[f"{key}_gap_pairs_weight_asymmetric"] = len(asym)
        out[f"{key}_gap_weight_sum_unordered"] = sum(pair_weight(p) for p in pairs)
        out[f"{key}_gap_rows_pre_lt_post"] = sum(r[0] < r[1] for r in gap)
        absent = set(groups["absent_from_neurons_json"])
        out[f"{key}_gap_pairs_touching_names_absent_from_neurons_json"] = sorted(
            _fmt_pair(p) for p in pairs if p & absent
        )

        if d["entries"] is not None:
            ce = [e for e in d["entries"] if e["typ"] == 0]
            ge = [e for e in d["entries"] if e["typ"] == 2]
            out[f"{key}_chem_unique_connector_ids"] = len({i for e in ce for i in e["ids"]})
            out[f"{key}_chem_syn_value_sum"] = sum(sum(e["syn"]) for e in ce)
            out[f"{key}_chem_entries_with_syn_value_gt1"] = sorted(
                f"{e['pre']}->{e['post']}:len{len(e['syn'])}:sum{sum(e['syn'])}" for e in ce if any(v > 1 for v in e["syn"])
            )
            out[f"{key}_chem_sum_unique_ids_per_edge"] = sum(len(set(e["ids"])) for e in ce)
            out[f"{key}_chem_entries_with_repeated_id"] = sum(len(set(e["ids"])) != len(e["ids"]) for e in ce)
            all_ids = [i for e in ce for i in e["ids"]]
            out[f"{key}_chem_connector_id_min"] = min(all_ids)
            out[f"{key}_chem_connector_id_max"] = max(all_ids)
            out[f"{key}_chem_records_minus_unique_ids_per_edge"] = sum(len(e["ids"]) - len(set(e["ids"])) for e in ce)
            mult = Counter()
            ids_3plus = []
            dup_records = []
            for e in ce:
                for i, n in Counter(e["ids"]).items():
                    mult[n] += 1
                    if n >= 3:
                        ids_3plus.append(f"{e['pre']}->{e['post']}:id{i}x{n}")
                recs = Counter(zip(e["ids"], e["pre_tid"], e["post_tid"]))
                dup_records += [
                    f"{e['pre']}->{e['post']}:id{i}:pre_tid{a}:post_tid{b}x{n}" for (i, a, b), n in recs.items() if n > 1
                ]
            out[f"{key}_chem_id_multiplicity_within_entry"] = {str(n): c for n, c in sorted(mult.items())}
            out[f"{key}_chem_ids_occurring_3plus_times_within_entry"] = sorted(ids_3plus)
            out[f"{key}_chem_identical_records_within_entry"] = sorted(dup_records)
            out[f"{key}_gap_unique_connector_ids"] = len({i for e in ge for i in e["ids"]})
            out[f"{key}_connector_ids_shared_by_chem_and_gap"] = len(
                {i for e in ce for i in e["ids"]} & {i for e in ge for i in e["ids"]}
            )

    comparisons = (
        ("wna_csv", "nemanode_859641e"),
        ("cect_xlsx", "nemanode_baf8ec8"),
        ("wna_csv", "cect_xlsx"),
        ("nemanode_859641e", "nemanode_baf8ec8"),
        ("nemanode_8a50236", "nemanode_859641e"),
    )
    for k in DATASETS:
        for a, b in comparisons:
            res = _compare(data[(a, k)]["rows"], data[(b, k)]["rows"])
            for rk, rv in res.items():
                out[f"cmp_d{k}_{a}__vs__{b}_{rk}"] = rv
        # The same NemaNode comparison with weight = sum of syn values instead
        # of len(syn), because the pre-revision files store a few synapses with
        # syn value 2.
        a, b = "nemanode_859641e", "nemanode_baf8ec8"
        res = _compare(_rows_sum_syn(data[(a, k)]["entries"]), _rows_sum_syn(data[(b, k)]["entries"]))
        for rk, rv in res.items():
            out[f"cmp_d{k}_{a}__vs__{b}_weight_sum_syn_{rk}"] = rv

    out.update(_paper_table_counts(raw, data, paper_neurons))
    return out


def _read_csv_dicts(path) -> list[dict]:
    import csv

    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def _read_paper_matrix(ws) -> tuple[list[tuple[str, str, str, int]], dict, tuple[set, set]]:
    """Read one sheet of synapse_count_matrices.xlsx as (pre, post, "chemical", count) rows.

    Layout, as written by the authors' export_matrices.py: row 3 holds the
    pre names from column 4 on; column 3 holds the post names from row 5 on;
    rows 1-2 and columns 1-2 hold the "Pre"/"Post" and cell-type labels. A
    type label is written only on the first row (column) of its block.
    Returns rows, facts, and the (pre names, post names) that define the
    matrix's cell set.
    """
    values = [list(r) for r in ws.iter_rows(values_only=True)]
    pre_names = values[2][3:]
    post_rows = [r for r in values[4:] if r[2] is not None]
    rows = []
    value_types = Counter()
    for r in post_rows:
        for pre, v in zip(pre_names, r[3:]):
            value_types[type(v).__name__] += 1
            if v not in (None, 0):
                rows.append((pre, r[2], "chemical", int(v)))
    post_by_label: dict = {}
    label = None
    for r in post_rows:
        label = r[1] if r[1] is not None else label
        post_by_label.setdefault(label, []).append(r[2])
    facts = {
        "pre_columns": len(pre_names),
        "post_rows": len(post_rows),
        "pre_type_labels": [v for v in values[1][3:] if v is not None],
        "post_type_labels": [r[1] for r in values[4:] if r[1] is not None],
        "post_rows_by_type_label": {k: len(v) for k, v in sorted(post_by_label.items())},
        "post_rows_labelled_Other": sorted(post_by_label.get("Other", [])),
        "cell_value_types": dict(sorted(value_types.items())),
        "nonzero": len(rows),
        "sum": sum(r[3] for r in rows),
    }
    return rows, facts, (set(pre_names), {r[2] for r in post_rows})


def _paper_table_counts(raw, data, paper_neurons) -> dict:
    """Counts from the first author's tables (see record_counts for definitions)."""
    import openpyxl

    out: dict = {}
    out["paper_neurons_csv_classes"] = len(paper_neurons)
    out["paper_neurons_csv_members_sum"] = sum(int(r["members"]) for r in paper_neurons)
    out["paper_neurons_csv_classes_by_type"] = dict(sorted(Counter(r["type"] for r in paper_neurons).items()))
    by_type = Counter()
    for r in paper_neurons:
        by_type[r["type"]] += int(r["members"])
    out["paper_neurons_csv_members_by_type"] = dict(sorted(by_type.items()))
    out["paper_neurons_csv_classes_by_integration"] = dict(sorted(Counter(r["integration"] for r in paper_neurons).items()))
    out["paper_neurons_csv_lists_CAN"] = any(r["class"] == "CAN" for r in paper_neurons)

    wb = openpyxl.load_workbook(raw / "synapse_count_matrices.xlsx", read_only=True)
    try:
        out["paper_matrices_sheetnames"] = list(wb.sheetnames)
        for k in DATASETS:
            rows, facts, (pre_set, post_set) = _read_paper_matrix(wb[f"Dataset{k}"])
            for fk, fv in sorted(facts.items()):
                out[f"paper_matrices_d{k}_{fk}"] = fv
            for other in ("nemanode_baf8ec8", "wna_csv", "nemanode_8a50236"):
                chem_only = [r for r in data[(other, k)]["rows"] if r[2] == "chemical"]
                res = _compare(rows, chem_only)
                for rk in ("chem_identical", "chem_edges_only_in_a", "chem_edges_only_in_b",
                           "chem_weight_differs_a_vs_b", "chem_synapse_sum_a_minus_b"):
                    out[f"cmp_d{k}_paper_matrix__vs__{other}_{rk}"] = res[rk]
                # Like for like: only the cells the matrix can hold (pre in its
                # columns and post in its rows). Edges outside are listed.
                inside = [r for r in chem_only if r[0] in pre_set and r[1] in post_set]
                pfx = f"cmp_d{k}_paper_matrix__vs__{other}_restricted_to_matrix_cells"
                out[f"{pfx}_b_edges_outside"] = sorted(
                    f"{r[0]}->{r[1]}:{r[3]}" for r in chem_only if not (r[0] in pre_set and r[1] in post_set)
                )
                out[f"{pfx}_b_chem_edges"] = len(inside)
                out[f"{pfx}_b_chem_synapse_sum"] = sum(r[3] for r in inside)
                res = _compare(rows, inside)
                for rk in ("chem_identical", "chem_edges_only_in_a", "chem_edges_only_in_b",
                           "chem_weight_differs_a_vs_b", "chem_synapse_sum_a_minus_b"):
                    out[f"{pfx}_{rk}"] = res[rk]
    finally:
        wb.close()

    syn = _read_csv_dicts(raw / "synapse_list.csv")
    out["paper_synapse_list_columns"] = list(syn[0].keys())
    out["paper_synapse_list_rows"] = len(syn)
    datasets = sorted({r["dataset"] for r in syn})
    out["paper_synapse_list_rows_by_dataset"] = {d: sum(r["dataset"] == d for r in syn) for d in datasets}
    out["paper_synapse_list_unique_syn_id_by_dataset"] = {
        d: len({r["syn_id"] for r in syn if r["dataset"] == d}) for d in datasets
    }
    out["paper_synapse_list_pre_post_pairs_by_dataset"] = {
        d: len({(r["pre"], r["post"]) for r in syn if r["dataset"] == d}) for d in datasets
    }
    out["paper_synapse_list_size_minus1_by_dataset"] = {
        d: sum(r["dataset"] == d and r["size"] == "-1" for r in syn) for d in datasets
    }
    for k in DATASETS:
        dk = [r for r in syn if r["dataset"] == f"Dataset{k}"]
        out[f"paper_synapse_list_d{k}_post_is_branch_values"] = dict(sorted(Counter(r["post_is_branch"] for r in dk).items()))
        counts = Counter((r["pre"], r["post"]) for r in dk)
        rows = [(a, b, "chemical", n) for (a, b), n in counts.items()]
        chem_only = [r for r in data[("nemanode_baf8ec8", k)]["rows"] if r[2] == "chemical"]
        res = _compare(rows, chem_only)
        for rk in ("chem_identical", "chem_edges_only_in_a", "chem_edges_only_in_b", "chem_weight_differs_a_vs_b"):
            out[f"cmp_d{k}_paper_synapse_list__vs__nemanode_baf8ec8_{rk}"] = res[rk]
        nn_ids = {str(i) for e in data[("nemanode_baf8ec8", k)]["entries"] if e["typ"] == 0 for i in e["ids"]}
        sl_ids = {r["syn_id"] for r in dk}
        out[f"cmp_d{k}_paper_synapse_list__vs__nemanode_baf8ec8_syn_id_sets_equal"] = nn_ids == sl_ids

    cls = _read_csv_dicts(raw / "connection_classifications.csv")
    out["paper_classifications_columns"] = list(cls[0].keys())
    out["paper_classifications_rows"] = len(cls)
    out["paper_classifications_duplicate_pairs"] = len(cls) - len({(r["pre"], r["post"]) for r in cls})
    out["paper_classifications_by_value"] = dict(sorted(Counter(r["classification"] for r in cls).items()))
    cls_np = [r for r in cls if r["classification"] != _POSTEMB_TABLE]
    classified = {(r["pre"], r["post"]) for r in cls_np}
    out["paper_classifications_rows_not_postembryonic"] = len(cls_np)
    out["paper_classifications_rows_not_postembryonic_self"] = sum(r["pre"] == r["post"] for r in cls_np)
    out["paper_classifications_pairs_not_postembryonic"] = len(classified)
    sl_pairs = {d: {(r["pre"], r["post"]) for r in syn if r["dataset"] == d} for d in datasets}
    present = {d: len(sl_pairs[d] & classified) for d in datasets}
    out["paper_classified_connections_present_by_dataset"] = present
    out["paper_classified_connections_present_sum_over_datasets"] = sum(present.values())
    out["paper_classified_connections_present_n_datasets"] = len(present)
    out["paper_classified_connections_present_mean_over_datasets"] = sum(present.values()) / len(present)

    out.update(_nemanode_history_counts(raw, syn, classified))
    out.update(_toolbox_validation_counts(raw, data))
    return out


def _nemanode_history_counts(raw, syn, table_classified: set) -> dict:
    """NemaNode head.annotations.json and pre-revision datasets 1-8 (see record_counts)."""
    out: dict = {}
    ann = json.loads((raw / "nemanode_859641e__head.annotations.json").read_text(encoding="utf-8"))
    for cat, pairs in ann.items():
        for p in pairs:
            if not (isinstance(p, list) and len(p) == 2 and all(isinstance(x, str) for x in p)):
                raise ValueError(f"head.annotations.json {cat}: unexpected entry {p!r}")
    pairs_by_cat = {cat: [tuple(p) for p in pairs] for cat, pairs in ann.items()}
    pfx = "nemanode_head_annotations"
    out[f"{pfx}_categories_entries"] = {c: len(v) for c, v in sorted(pairs_by_cat.items())}
    out[f"{pfx}_entries_total"] = sum(len(v) for v in pairs_by_cat.values())
    cat_count = Counter(p for v in pairs_by_cat.values() for p in set(v))
    out[f"{pfx}_distinct_pairs"] = len(cat_count)
    out[f"{pfx}_duplicate_entries_within_category"] = sum(len(v) - len(set(v)) for v in pairs_by_cat.values())
    out[f"{pfx}_pairs_in_more_than_one_category"] = sum(1 for n in cat_count.values() if n > 1)
    out[f"{pfx}_self_pairs"] = sum(a == b for a, b in cat_count)
    ann_np = {p for c, v in pairs_by_cat.items() if c != _POSTEMB_NEMANODE for p in v}
    out[f"{pfx}_pairs_not_postembryonic"] = len(ann_np)
    out[f"cmp_{pfx}__vs__paper_classifications_not_postembryonic_pairs_only_in_annotations"] = sorted(
        f"{a}->{b}" for a, b in ann_np - table_classified
    )
    out[f"cmp_{pfx}__vs__paper_classifications_not_postembryonic_pairs_only_in_table"] = sorted(
        f"{a}->{b}" for a, b in table_classified - ann_np
    )

    pre = {k: _nemanode_chem(raw / f"nemanode_859641e__witvliet_2020_{k}.json") for k in range(1, 9)}
    out["nemanode_859641e_by_dataset_chem_edges"] = {str(k): len(v) for k, v in pre.items()}
    out["nemanode_859641e_by_dataset_chem_synapse_sum"] = {str(k): sum(v.values()) for k, v in pre.items()}
    present = {str(k): len(set(v) & ann_np) for k, v in pre.items()}
    q = f"{pfx}_not_postembryonic_present_in_nemanode_859641e_chem"
    out[f"{q}_by_dataset"] = present
    out[f"{q}_sum_over_datasets"] = sum(present.values())
    out[f"{q}_n_datasets"] = len(present)
    out[f"{q}_mean_over_datasets"] = sum(present.values()) / len(present)
    sl = {}
    for r in syn:
        sl.setdefault(r["dataset"], Counter())[(r["pre"], r["post"])] += 1
    present = {d: len(set(c) & ann_np) for d, c in sorted(sl.items())}
    q = f"{pfx}_not_postembryonic_present_in_paper_synapse_list"
    out[f"{q}_by_dataset"] = present
    out[f"{q}_sum_over_datasets"] = sum(present.values())
    out[f"{q}_n_datasets"] = len(present)
    out[f"{q}_mean_over_datasets"] = sum(present.values()) / len(present)

    # Dataset-1 chemical connections followed into the adult datasets: how many
    # are still present, and the adult synapse sum over them.
    for label, d1, adults in (
        ("nemanode_859641e", pre[1], {k: pre[k] for k in DATASETS}),
        ("paper_synapse_list", sl["Dataset1"], {k: sl[f"Dataset{k}"] for k in DATASETS}),
    ):
        q = f"d1_edges_in_adult__{label}"
        out[f"{q}_d1_chem_edges"] = len(d1)
        out[f"{q}_d1_chem_synapse_sum"] = sum(d1.values())
        for k, adult in adults.items():
            kept = [p for p in d1 if p in adult]
            out[f"{q}_d{k}_present"] = len(kept)
            out[f"{q}_d{k}_synapse_sum_on_present"] = sum(adult[p] for p in kept)
            out[f"{q}_d{k}_mean_on_present"] = sum(adult[p] for p in kept) / len(kept)
            out[f"{q}_d{k}_mean_on_all_d1_edges"] = sum(adult[p] for p in kept) / len(d1)
    return out


def _read_toolbox_expected_yaml(path) -> dict:
    """Parse a ConnectomeToolbox ``*_expected_data.yaml`` of the fixed layout it uses; anything else raises."""
    doc = {"reader": None, "connection_lists": []}
    cur = conn = None
    for n, line in enumerate(path.read_text(encoding="utf-8").split("\n"), start=1):
        if line == "" or line == "connection_lists:":
            continue
        if line.startswith("reader: "):
            doc["reader"] = line[len("reader: "):]
        elif line == "- connections:":
            cur = {"connections": []}
            doc["connection_lists"].append(cur)
        elif line.startswith("  - pre: ") and cur is not None:
            conn = {"pre": line[len("  - pre: "):]}
            cur["connections"].append(conn)
        elif line.startswith("    post: ") and conn is not None:
            conn["post"] = line[len("    post: "):]
        elif line.startswith("    weight: ") and conn is not None:
            conn["weight"] = int(line[len("    weight: "):])
        elif line.startswith("  synapse: ") and cur is not None:
            cur["synapse"] = line[len("  synapse: "):]
        elif line.startswith("  total_nonzero_conns: ") and cur is not None:
            cur["total_nonzero_conns"] = int(line[len("  total_nonzero_conns: "):])
        else:
            raise ValueError(f"{path.name} line {n}: unexpected {line!r}")
    return doc


def _toolbox_validation_counts(raw, data) -> dict:
    """ConnectomeToolbox expected values for datasets 7 and 8, compared with every version (see record_counts)."""
    import openpyxl

    out: dict = {}
    md = (raw / "WitvlietEtAl2021.md").read_text(encoding="utf-8")
    out["toolbox_validation_md_names_MOESM4"] = "41586_2021_3778_MOESM4_ESM.xlsx" in md
    out["toolbox_validation_md_says_chemical_weights_read_from_supplement"] = (
        "The chemical connection weights below were read from the supplementary information spreadsheet" in md
    )
    summary = (raw / "Witvliet_2021.md").read_text(encoding="utf-8")
    out["toolbox_docs_paraphrase_lines_with_3113_or_1292"] = sorted(
        line.strip() for line in summary.split("\n") if "3,113" in line or "1,292" in line
    )
    wb = openpyxl.load_workbook(raw / "synapse_count_matrices.xlsx", read_only=True)
    try:
        matrix = {k: _edge_dicts(_read_paper_matrix(wb[f"Dataset{k}"])[0])[0] for k in DATASETS}
    finally:
        wb.close()
    for k in DATASETS:
        doc = _read_toolbox_expected_yaml(raw / f"WitvlietDataReader{k}_expected_data.yaml")
        pfx = f"toolbox_validation_d{k}"
        out[f"{pfx}_reader"] = doc["reader"]
        lists = {cl["synapse"]: cl for cl in doc["connection_lists"]}
        if len(lists) != len(doc["connection_lists"]):
            raise ValueError(f"{pfx}: a synapse class is listed twice")
        for syncls, cl in sorted(lists.items()):
            out[f"{pfx}_{syncls}_total_nonzero_conns"] = cl.get("total_nonzero_conns")
            out[f"{pfx}_{syncls}_sample_weights"] = [f"{c['pre']}->{c['post']}:{c['weight']}" for c in cl["connections"]]
        sample = lists["Generic_CS"]["connections"]
        stores = {v: _edge_dicts(data[(v, k)]["rows"])[0] for v in VERSIONS}
        stores["paper_matrix"] = matrix[k]
        out[f"{pfx}_Generic_CS_sample_mismatches_by_version"] = {
            v: [f"{c['pre']}->{c['post']}:{c['weight']}vs{chem.get((c['pre'], c['post']))}"
                for c in sample if chem.get((c["pre"], c["post"])) != c["weight"]]
            for v, chem in sorted(stores.items())
        }
        x = f"cect_xlsx_d{k}"
        xchem, xgap = _edge_dicts(data[("cect_xlsx", k)]["rows"])
        self_pairs = sum(len(p) == 1 for p in xgap)
        out[f"{x}_chem_nonzero_cells"] = len(xchem)
        out[f"{x}_gap_nonzero_cells_symmetric"] = 2 * (len(xgap) - self_pairs) + self_pairs
    return out


def _rows_sum_syn(entries):
    return [(e["pre"], e["post"], _NEMANODE_TYP.get(e["typ"], f"typ={e['typ']}"), sum(e["syn"])) for e in entries]
